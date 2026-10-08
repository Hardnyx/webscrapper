"""Shared transport, period planning and cache contract for SBS monthly Excel."""
from datetime import date
from io import BytesIO
import re
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
import pandas as pd

from fuentes_financieras.exceptions import InvalidQueryError, PeriodUnavailableError, SchemaChangedError
from fuentes_financieras.models import FetchResult, PeriodRequest
from fuentes_financieras.provider import DatasetProvider
from fuentes_financieras.storage import utc_now_iso
from fuentes_financieras.transports.curl_chrome import CurlChromeTransport

INDEX = 'https://www.sbs.gob.pe/app/stats_net/stats/EstadisticaSistemaFinancieroResultados.aspx?c='
MONTHS = dict(zip(('en', 'fe', 'ma', 'ab', 'my', 'jn', 'jl', 'ag', 'se', 'oc', 'no', 'di'), range(1, 13)))


def clean(value):
    return '' if pd.isna(value) else ' '.join(str(value).split())


def month(value):
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}', value):
        raise InvalidQueryError('Use períodos YYYY-MM.')
    try:
        target = date.fromisoformat(value + '-01')
    except ValueError as exc:
        raise InvalidQueryError('Mes inválido.') from exc
    if target < date(2013, 1, 1) or target > date.today():
        raise InvalidQueryError('Período admitido: enero de 2013 hasta el mes actual.')
    return target


def discover_files(html, code):
    """Use published links, never synthesize download URLs."""
    found = {}
    for anchor in BeautifulSoup(html, 'html.parser').find_all('a', href=True):
        url = urljoin(INDEX + code, anchor['href'])
        parsed = urlparse(url)
        match = re.search(r'/' + code + r'-(\w{2})(\d{4})\.xlsx?$', parsed.path, re.I)
        if not match:
            continue
        if parsed.scheme != 'https' or parsed.hostname != 'intranet2.sbs.gob.pe':
            raise SchemaChangedError('Host de descarga SBS inesperado.')
        token, year = match.groups()
        if token.lower() not in MONTHS:
            raise SchemaChangedError('Mes del enlace SBS desconocido.')
        period = f'{year}-{MONTHS[token.lower()]:02d}'
        if period in found and found[period] != url:
            raise SchemaChangedError(f'Enlaces contradictorios para {period}.')
        found[period] = url
    if not found:
        raise SchemaChangedError('El índice no contiene enlaces de publicaciones mensuales.')
    return found


def read_workbook(content):
    if content.startswith(b'PK\x03\x04'):
        engine = 'openpyxl'
    elif content.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        engine = 'xlrd'
    else:
        raise SchemaChangedError('La descarga no es un libro Excel XLS/XLSX.')
    try:
        return pd.read_excel(BytesIO(content), sheet_name=None, header=None, engine=engine)
    except Exception as exc:
        raise SchemaChangedError('Libro Excel ilegible.') from exc


def cell_formats(content, positions):
    """Read number formats for explicit zero-based sheet/row/column positions."""
    if content.startswith(b'PK'):
        import openpyxl
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            return {key: book[key[0]].cell(key[1]+1, key[2]+1).number_format for key in positions}
        finally:
            book.close()
    import xlrd
    book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    return {key: book.format_map[book.xf_list[book.sheet_by_name(key[0]).cell(key[1], key[2]).xf_index].format_key].format_str
            for key in positions}


class MonthlyExcelProvider(DatasetProvider):
    codes = {}

    def __init__(self, spec):
        super().__init__(spec)
        self.transport = CurlChromeTransport(state_dir=self.storage.state_root, timeout=30)
        self.indices = {}

    def single_request(self, *, periodo=None, tipo='B', **query):
        if query or tipo not in self.codes:
            raise InvalidQueryError('Use periodo=YYYY-MM y tipo=B/F/C/R.')
        target = month(periodo)
        return PeriodRequest(f'{tipo}:{periodo}', f'tipo={tipo}/year={target.year}', {'periodo': periodo, 'tipo': tipo}, mutable=True)

    def plan_sync(self, *, desde, hasta=None, tipos=('B', 'F', 'C', 'R'), **query):
        start, end = month(desde), month(hasta or desde)
        selected = [tipos] if isinstance(tipos, str) else list(tipos)
        if query or not selected or len(set(selected)) != len(selected) or set(selected) - self.codes.keys() or end < start:
            raise InvalidQueryError('Rango o tipos inválidos.')
        while start <= end:
            for tipo in selected:
                yield self.single_request(periodo=start.strftime('%Y-%m'), tipo=tipo)
            start = date(start.year + (start.month == 12), start.month % 12 + 1, 1)

    def _fetch_period(self, request):
        tipo, period = request.params['tipo'], request.params['periodo']
        if tipo not in self.indices:
            self.indices[tipo] = discover_files(self.transport.request('GET', INDEX + self.codes[tipo]).text, self.codes[tipo])
        if period not in self.indices[tipo]:
            raise PeriodUnavailableError(f'SBS no enlaza la publicación de {tipo} para {period}.')
        url = self.indices[tipo][period]
        content = self.transport.request('GET', url).content
        data, notes = self.parse_workbook(content, entity_type=tipo, period=period, source_url=url, retrieved_at=utc_now_iso())
        return FetchResult(self.spec.dataset_id, data, {'period': period, 'entity_type': tipo, 'source_url': url, 'notes': notes}, content)

    def sync(self, **query):
        self.indices = {}
        return super().sync(**query)

    def filter_loaded(self, data, *, desde=None, hasta=None, tipos=None):
        for value, upper in ((desde, False), (hasta, True)):
            if value is not None:
                month(value)
                data = data[data.period <= value] if upper else data[data.period >= value]
        if tipos is not None:
            selected = [tipos] if isinstance(tipos, str) else list(tipos)
            if set(selected) - self.codes.keys():
                raise InvalidQueryError('Tipos admitidos: B/F/C/R.')
            data = data[data.entity_type.isin(selected)]
        return data.drop(columns=['_period_key'], errors='ignore')
