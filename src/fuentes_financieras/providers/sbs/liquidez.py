"""Independent published liquidity, coverage and stable funding observations."""
from calendar import monthrange
from datetime import date, datetime
from io import BytesIO
import math
import re

import pandas as pd

from fuentes_financieras.exceptions import SchemaChangedError
from ._excel_mensual import MonthlyExcelProvider, clean, month, read_workbook

LIQUIDITY_CODES = dict(zip('BFCR', ('B-2340', 'B-3250', 'C-1244', 'C-2249')))
COVERAGE_CODES = dict(zip('BFCR', ('B-230809', 'B-230810', 'B-230811', 'B-230812')))
FUNDING_CODES = dict(zip('BFCR', ('B-234021', 'B-230213', 'C-120212', 'B-230820')))
MONTH_NAMES = dict(zip(('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
    'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre'), range(1, 13)))


def normalized(value):
    return clean(value).replace('\\n', ' ').lower()


def numeric(value):
    if pd.isna(value):
        return float('nan')
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise SchemaChangedError('Valor de liquidez no numérico.')
    return float(value)


def closing(period):
    target = month(period)
    return date(target.year, target.month, monthrange(target.year, target.month)[1])


def number_formats(content, row, columns):
    """Read the actual percentage format; never infer a scale from magnitude."""
    if content.startswith(b'PK'):
        import openpyxl
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            return {(s.title, col): s.cell(row+1, col+1).number_format for s in book for col in columns}
        finally:
            book.close()
    import xlrd
    book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    return {(s.name, col): book.format_map[book.xf_list[s.cell(row, col).xf_index].format_key].format_str
            for s in book.sheets() for col in columns}


def finish(rows, notes):
    if not rows:
        raise SchemaChangedError('Cuadro sin observaciones de liquidez.')
    data = pd.DataFrame(rows)
    for col in data:
        data[col] = data[col].astype('float64' if col in ('value', 'source_value', 'unit_multiplier')
                                    else ('int64' if col == 'source_row' else 'string'))
    return data, notes


def record(*, entity_type, period, name, sheet, row, metric, label, value, unit,
           currency, basis, start, end, source_url, retrieved_at, source_value=None, flags=''):
    return dict(period=period, period_date=end.isoformat(), observation_start=start.isoformat(),
        observation_end=end.isoformat(), frequency='quarterly' if basis == 'daily_average_over_quarter' else 'monthly',
        entity_type=entity_type, entity_name=name,
        entity_scope='system_aggregate' if name.upper().startswith(('TOTAL ', 'CONSOLIDADO')) else 'entity',
        metric=metric, source_metric=label, value=value, source_value=value if source_value is None else source_value,
        unit=unit, unit_multiplier=1000.0 if unit.startswith('thousands_') else float('nan'),
        currency=currency, measurement_basis=basis, data_quality_flags=flags,
        source_sheet=str(sheet), source_row=row+1, source='SBS', source_url=source_url, retrieved_at=retrieved_at)


def parse_liquidity(content, *, entity_type, period, source_url, retrieved_at):
    book = read_workbook(content)
    candidates = [(s, d) for s, d in book.items() if not d.empty and
        any(normalized(v).startswith('ratios de liquidez en moneda nacional') for v in d.iloc[:3, 0])]
    if len(candidates) != 1:
        raise SchemaChangedError('Hoja de liquidez ausente o ambigua.')
    sheet, frame = candidates[0]
    if frame.shape[1] != 8 or len(frame) < 8:
        raise SchemaChangedError('Dimensiones de liquidez no reconocidas.')
    heads = [i for i, v in frame.iloc[:7, 0].items() if clean(v) == 'Empresas']
    if len(heads) != 1:
        raise SchemaChangedError('Encabezado de empresas ausente o ambiguo.')
    head = heads[0]
    dates = [v for v in frame.iloc[:head, 0] if isinstance(v, (datetime, pd.Timestamp))]
    end = closing(period)
    if len(dates) != 1 or dates[0].date() != end:
        raise SchemaChangedError('Fecha de liquidez distinta del mes solicitado.')
    for first, currency_word, unit_word in ((1, 'nacional', 'soles'), (5, 'extranjera', 'dólares')):
        if not normalized(frame.iloc[head, first]).startswith('liquidez en moneda ' + currency_word):
            raise SchemaChangedError('Grupo monetario de liquidez cambiado.')
        for col, label in ((first, 'activos líquidos'), (first+1, 'pasivos de'), (first+2, 'ratio de liquidez')):
            text = normalized(frame.iloc[head+1, col])
            if not text.startswith(label):
                raise SchemaChangedError('Encabezado de indicador de liquidez cambiado.')
            combined = normalized(frame.iloc[head, first]) + ' ' + text
            if ('porcentaje' if col == first+2 else 'miles de ' + unit_word) not in combined:
                raise SchemaChangedError('Unidad de liquidez ausente o cambiada.')
    stop = next((i for i in range(head+2, len(frame)) if normalized(frame.iloc[i, 0]).startswith(('fuente:', 'nota:'))), None)
    if stop is None:
        raise SchemaChangedError('Falta el cierre de fuente de liquidez.')
    notes = [{'sheet': sheet, 'row': i+1, 'text': clean(frame.iloc[i, 0])}
             for i in range(stop, len(frame)) if clean(frame.iloc[i, 0])]
    rows, names = [], set()
    for row in range(head+2, stop):
        name = clean(frame.iloc[row, 0])
        if not name:
            if frame.iloc[row, [1, 2, 3, 5, 6, 7]].notna().any():
                raise SchemaChangedError('Valores de liquidez sin entidad.')
            continue
        if name in names:
            raise SchemaChangedError('Entidad de liquidez duplicada.')
        names.add(name)
        for first, currency in ((1, 'PEN'), (5, 'USD')):
            values = [numeric(frame.iloc[row, col]) for col in range(first, first+3)]
            mismatch = all(math.isfinite(v) for v in values) and values[1] > 0 and abs(values[0]/values[1]*100-values[2]) > .02
            for offset, metric in enumerate(('liquid_assets', 'short_term_liabilities', 'liquidity_ratio')):
                value = values[offset]
                flags = ['source_value_missing'] if math.isnan(value) else []
                if mismatch:
                    flags.append('published_ratio_mismatch')
                rows.append(record(entity_type=entity_type, period=period, name=name, sheet=sheet,
                    row=row, metric=metric, label=clean(frame.iloc[head+1, first+offset]), value=value,
                    unit='percent' if offset == 2 else 'thousands_' + currency, currency=currency,
                    basis='month_end', start=month(period), end=end, source_url=source_url,
                    retrieved_at=retrieved_at, flags=';'.join(flags)))
    if not any(n.upper().startswith('TOTAL ') for n in names):
        raise SchemaChangedError('Falta el agregado de liquidez.')
    return finish(rows, notes)


def parse_disclosures(content, *, entity_type, period, source_url, retrieved_at, kind):
    book = read_workbook(content)
    formats = number_formats(content, 54, [7]) if kind == 'funding' else number_formats(content, 34, [5, 7, 9])
    rows, notes, names = [], [], set()
    end_index = closing(period)
    for sheet, frame in book.items():
        if frame.shape[1] != (10 if kind == 'coverage' else 8) or len(frame) != (40 if kind == 'coverage' else 61):
            raise SchemaChangedError('Hoja de divulgación de liquidez con estructura desconocida.')
        title = 'Ratio de Cobertura de Liquidez' if kind == 'coverage' else 'Ratio de Financiación Neta Estable'
        if clean(frame.iloc[1, 1]) != title:
            raise SchemaChangedError('Título de divulgación de liquidez cambiado.')
        flags = []
        if kind == 'coverage':
            text = clean(frame.iloc[2, 1])
            match = re.fullmatch(r'Saldos y ratio promedio diario de (\w+) a (\w+) de (\d{4}) \(1\)', text)
            if not match or any(m not in MONTH_NAMES for m in match.groups()[:2]):
                raise SchemaChangedError('Período trimestral del RCL no reconocido.')
            first, last, year = match.groups(); a, b = MONTH_NAMES[first], MONTH_NAMES[last]
            if a not in (1, 4, 7, 10) or b != a+2:
                raise SchemaChangedError('Trimestre RCL no reconocido.')
            start = date(int(year), a, 1); end = closing(f'{year}-{b:02d}')
            if end > end_index:
                raise SchemaChangedError('El trimestre RCL es posterior al mes del enlace.')
            if end != end_index:
                flags.append('source_period_differs_from_index')
            name = clean(frame.iloc[5, 4]); stop = 35
            for col, group in ((4, 'MONEDA NACIONAL (En miles de soles)'), (6, 'MONEDA EXTRANJERA (En miles de dólares)'), (8, 'TOTAL (En miles de soles)')):
                if clean(frame.iloc[6, col]) != group or clean(frame.iloc[7, col]) != 'Importe Base (promedio)' or clean(frame.iloc[7, col+1]) != 'Importe Ajustado (promedio)':
                    raise SchemaChangedError('Unidades o bases del RCL cambiadas.')
            if any('%' in formats[sheet, col] for col in (5, 7, 9)):
                raise SchemaChangedError('Formato del RCL cambiado: requiere revisión de escala.')
            specs = [(31, 'high_quality_liquid_assets', 'Total ALAC'), (32, 'inflows_30_days', 'Total Flujos Entrantes 30 días'),
                     (33, 'outflows_30_days', 'Total Flujos Salientes 30 días'), (34, 'liquidity_coverage_ratio', 'RATIO DE COBERTURA DE LIQUIDEZ (%) (2)')]
        else:
            observed = frame.iloc[2, 1]
            if not isinstance(observed, (datetime, pd.Timestamp)) or observed.date() != end_index:
                raise SchemaChangedError('Fecha del RFNE distinta del mes solicitado.')
            start, end = month(period), end_index
            name = clean(frame.iloc[4, 3]); stop = 55
            if clean(frame.iloc[5, 7]) != 'VALOR PONDERADO (En miles de Soles)' or clean(frame.iloc[5, 3]) != 'VALOR NO PONDERADO POR VENCIMIENTO RESIDUAL (En miles de Soles)':
                raise SchemaChangedError('Unidades del RFNE cambiadas.')
            if formats[sheet, 7] not in ('0%', '0.0%', '0.00%'):
                raise SchemaChangedError('Formato porcentual del RFNE no reconocido.')
            specs = [(52, 'available_stable_funding', 'Total Financiación Estable Disponible'),
                (53, 'required_stable_funding', 'Total Financiación Estable Requerida'),
                (54, 'net_stable_funding_ratio', 'RATIO DE FINANCIACIÓN NETA ESTABLE (%)')]
        if kind == 'funding':
            available, required, ratio = [numeric(frame.iloc[i, 7]) for i in (52, 53, 54)]
            if all(math.isfinite(v) for v in (available, required, ratio)) and required > 0 and abs(available/required-ratio) > .0002:
                flags.append('published_ratio_mismatch')
        if not name or name in names:
            raise SchemaChangedError('Entidad de divulgación ausente o duplicada.')
        names.add(name)
        if not clean(frame.iloc[stop, 1]).startswith('Fuente:'):
            raise SchemaChangedError('Fuente de divulgación ausente.')
        notes.append({'sheet': sheet, 'index_period': period, 'observation_start': start.isoformat(),
            'observation_end': end.isoformat(), 'texts': [clean(frame.iloc[i, 1])+' '+clean(frame.iloc[i, 2])
                for i in range(stop, len(frame)) if clean(frame.iloc[i, 1]) or clean(frame.iloc[i, 2])]})
        for row, metric, label in specs:
            if clean(frame.iloc[row, 2]) != label:
                raise SchemaChangedError('Indicador de divulgación cambiado.')
            ratio = metric.endswith('_ratio')
            columns = [(7, 'PEN', 'weighted')] if kind == 'funding' else [
                (col+offset, currency, basis) for col, currency in ((4, 'PEN'), (6, 'USD'), (8, 'TOTAL_PEN'))
                for offset, basis in ((0, 'base'), (1, 'adjusted')) if not ratio or offset == 1]
            for col, currency, basis in columns:
                original = numeric(frame.iloc[row, col]); value = original*100 if kind == 'funding' and ratio else original
                cell_flags = flags + (['source_value_missing'] if math.isnan(original) else [])
                rows.append(record(entity_type=entity_type, period=period, name=name, sheet=sheet, row=row,
                    metric=metric, label=label, value=value, source_value=original,
                    unit='percent' if ratio else 'thousands_' + ('USD' if currency == 'USD' else 'PEN'),
                    currency=currency, basis='daily_average_over_quarter' if kind == 'coverage' else 'month_end',
                    start=start, end=end, source_url=source_url, retrieved_at=retrieved_at, flags=';'.join(cell_flags)) | {
                        'amount_basis': 'daily_ratio_average' if kind == 'coverage' and ratio else basis,
                        'source_number_format': formats[sheet, col] if ratio else ''})
    if not any(n.upper().startswith('CONSOLIDADO') for n in names):
        raise SchemaChangedError('Falta el consolidado de divulgación.')
    return finish(rows, notes)


def parse_coverage(content, **kwargs):
    return parse_disclosures(content, kind='coverage', **kwargs)


def parse_funding(content, **kwargs):
    return parse_disclosures(content, kind='funding', **kwargs)


class LiquidityProvider(MonthlyExcelProvider):
    codes = LIQUIDITY_CODES
    parser_version = '2026-10-08.1'
    parse_workbook = staticmethod(parse_liquidity)


class LiquidityCoverageProvider(MonthlyExcelProvider):
    codes = COVERAGE_CODES
    parser_version = '2026-10-08.1'
    parse_workbook = staticmethod(parse_coverage)


class StableFundingProvider(MonthlyExcelProvider):
    codes = FUNDING_CODES
    parser_version = '2026-10-08.1'
    parse_workbook = staticmethod(parse_funding)
