"""Synchronize SBS passive rates using the shared provider."""
from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from fuentes_financieras import source


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Tasas pasivas SBS B/C/F/R, con caché local.")
    parser.add_argument('--tipos', nargs='+', choices=['B', 'C', 'F', 'R'], default=['B', 'C', 'F', 'R'])
    parser.add_argument('--desde', required=True, help='YYYY-MM-DD')
    parser.add_argument('--hasta', default=datetime.now(ZoneInfo('America/Lima')).date().isoformat())
    parser.add_argument('--data-root', type=Path)
    parser.add_argument('--force', action='store_true')
    parser.add_argument('--keep-raw', action='store_true')
    parser.add_argument('--load-only', action='store_true', help='Leer solo el almacén local, sin consultas.')
    parser.add_argument('--excel', type=Path, help='Exportación opcional con tabla y filtros.')
    args = parser.parse_args(argv)
    try:
        start = datetime.strptime(args.desde, '%Y-%m-%d').date()
        end = datetime.strptime(args.hasta, '%Y-%m-%d').date()
    except ValueError:
        parser.error('Las fechas deben ser válidas y tener formato YYYY-MM-DD.')
    if end < start:
        parser.error('--hasta debe ser igual o posterior a --desde.')
    args.tipos = list(dict.fromkeys(args.tipos))
    return args


def export_excel(frame, path: Path):
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    import pandas as pd

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        frame.to_excel(writer, sheet_name='Tasas pasivas', index=False)
        sheet = writer.sheets['Tasas pasivas']
        sheet.freeze_panes = 'A2'
        table = Table(displayName='TasasPasivas', ref=f'A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        sheet.add_table(table)


def main(argv=None) -> int:
    import os
    args = parse_args(argv)
    if args.data_root:
        os.environ['FINANCIAL_SOURCES_DATA_ROOT'] = str(args.data_root.expanduser().resolve())
        # The registry caches providers; a changed root needs a fresh instance.
        from fuentes_financieras.registry import get_provider
        get_provider.cache_clear()
    dataset = source('pe.sbs.tasas_pasivas')
    query = dict(tipos=args.tipos, desde=args.desde, hasta=args.hasta)
    if not args.load_only:
        result = dataset.sync(**query, force=args.force, keep_raw=args.keep_raw)
        print(json.dumps({
            'solicitados': result.requested, 'descargados': result.downloaded,
            'en_cache': result.skipped_existing, 'sin_cambios': result.unchanged,
            'no_disponibles': result.unavailable, 'fallidos': result.failed,
            'detalle': result.details,
        }, ensure_ascii=False, default=str))
        if result.failed or result.unavailable:
            return 1
    frame = dataset.load(**query)
    print(f'Filas locales: {len(frame)}')
    if frame.empty:
        print('No hay datos para el período y tipos seleccionados.')
        return 1
    if args.excel:
        export_excel(frame, args.excel)
        print(f'Excel: {args.excel.resolve()}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
