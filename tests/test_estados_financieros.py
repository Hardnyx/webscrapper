"""Guard period, units, totals and cache behavior without network fixtures."""
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pandas as pd
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError, InvalidQueryError
from fuentes_financieras.providers.sbs.estados_financieros import (
    FinancialStatementsProvider, discover_files, parse_workbook,
)


def workbook(*, unit='(En miles de soles)', assets=100, period='2026-08-31', income=-5):
    book = Workbook()
    balance = book.active
    balance.title = 'balance'
    results = book.create_sheet('results')
    def block(sheet, start, title, section, accounts):
        sheet.cell(start, 1, title)
        sheet.cell(start + 1, 1, pd.Timestamp(period).to_pydatetime())
        sheet.cell(start + 2, 1, unit)
        sheet.cell(start + 4, 1, section)
        sheet.cell(start + 4, 2, 'Entidad*')
        for col, currency in enumerate(('MN', 'ME', 'TOTAL'), 2):
            sheet.cell(start + 5, col, currency)
        for row, (label, value) in enumerate(accounts, start + 7):
            sheet.cell(row, 1, label)
            # Foreign-currency amounts are expressed in soles too.
            if value is not None:
                for col, amount in ((2, value), (3, 0), (4, value)):
                    sheet.cell(row, col, amount)
    block(balance, 2, 'Balance General por Empresa', 'Activo', [('TOTAL ACTIVO', assets), ('   Caja', None)])
    block(balance, 16, 'Balance General por Empresa', 'Pasivo', [('TOTAL PASIVO', 80), ('PATRIMONIO', 20), ('TOTAL PASIVO Y PATRIMONIO', 100), ('Resultado Neto del Ejercicio', -5)])
    block(results, 2, 'Estado de Ganancias y Pérdidas por Empresa', '', [('RESULTADO NETO DEL EJERCICIO', income), ('   Otros', 0)])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def parse(content):
    return parse_workbook(content, entity_type='F', period='2026-08', source_url='https://source', retrieved_at='test')[0]


def test_units_signed_amounts_and_accumulation():
    data = parse(workbook())
    assert data[data.account == 'Caja'].amount.isna().all()
    assert len(data[data.account == 'Caja']) == 3
    assert set(data.unit) == {'thousands_PEN'}
    assert set(data.unit_multiplier) == {1000}
    assert set(data.entity_name) == {'Entidad'}
    assert set(data.source_entity_name) == {'Entidad*'}
    assert data[(data.statement == 'income') & (data.currency == 'TOTAL')].amount.tolist() == [-5, 0]
    assert set(data[data.statement == 'income'].measurement_basis) == {'year_to_date'}
    assert set(data[data.statement == 'balance'].measurement_basis) == {'closing_balance'}


@pytest.mark.parametrize('kwargs', [
    {'unit': '(En millones de soles)'}, {'period': '2026-07-31'},
    {'assets': 101}, {'income': -6},
])
def test_rejects_wrong_units_period_or_accounting_identity(kwargs):
    with pytest.raises(SchemaChangedError):
        parse(workbook(**kwargs))


def test_rejects_html_masquerading_as_excel():
    with pytest.raises(SchemaChangedError):
        parse(b'<html>Access denied</html>')


def test_published_link_discovery_and_guard():
    url = 'https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-3101-ag2026.XLS'
    assert discover_files(f'<a href="{url}">Agosto</a>', 'F') == {'2026-08': url}
    with pytest.raises(SchemaChangedError):
        discover_files('<html>Error de servidor</html>', 'F')
    with pytest.raises(SchemaChangedError):
        discover_files(f'<a href="{url.replace("intranet2.sbs.gob.pe", "example.org")}">Agosto</a>', 'F')


def test_cache_and_query_validation(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path))
    provider = FinancialStatementsProvider(CATALOG['pe.sbs.estados_financieros'])
    url = 'https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-3101-ag2026.XLS'
    calls = []
    def request(method, target):
        calls.append(target)
        return SimpleNamespace(text=f'<a href="{url}">Agosto</a>', content=workbook())
    monkeypatch.setattr(provider.transport, 'request', request)
    first = provider.sync(desde='2026-08', tipos=['F'], keep_raw=True)
    assert first.downloaded == 1 and not first.failed
    second = provider.sync(desde='2026-08', tipos=['F'])
    assert second.skipped_existing == 1 and len(calls) == 2
    assert len(provider.load(desde='2026-08', hasta='2026-08', tipos=['F'])) == 24
    assert provider.sync(desde='2026-07', tipos=['F']).unavailable == 1
    for query in ({'periodo': '2026-13'}, {'periodo': '2012-12'}, {'periodo': '2026-08', 'tipo': 'X'}):
        with pytest.raises(InvalidQueryError):
            provider.single_request(**query)
