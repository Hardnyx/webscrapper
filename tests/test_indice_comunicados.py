"""Validate discovery identities and fail closed before fetching PDFs."""
from types import SimpleNamespace
import pytest
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import InvalidQueryError,SchemaChangedError
from fuentes_financieras.providers.clasificadoras.indice_comunicados import (
    INDEX,action_reference,parse_index,parse_reference,ActionIndexProvider,ActionReferencesProvider,
)
URL='https://moodyslocal.com.pe/?post_type=rating-action&p=123'
FRIENDLY='https://moodyslocal.com.pe/reporte/rating-action/ejemplo/'
PDF='https://moodyslocal.com.pe/wp-content/uploads/2026/07/ejemplo.pdf'
TITLE='Moody’s Local Perú retira las clasificaciones de Ejemplo'

def index():
    return f'<html><table id="table_1" data-wpdatatable_id="35"><thead><tr><th>Fecha</th><th>Título</th></tr></thead><tbody><tr><td>30/07/2026</td><td><a href="{URL}">{TITLE}</a></td></tr></tbody></table></html>'

def page():
    return f'<html><head><link rel="canonical" href="{FRIENDLY}"><link rel="shortlink" href="https://moodyslocal.com.pe/?p=123"></head><body><div class="et_pb_title_container"><h1 class="entry-title">{TITLE}</h1><p><span class="published">Jul 30, 2026</span></p></div><a href="{PDF}">Download</a><span class="published">Otro artículo</span></body></html>'

def test_delivered_index_and_reference_only_do_not_establish_events():
    d,m=parse_index(index(),retrieved_at='checked')
    assert len(d)==1 and d.listed_date.iloc[0]=='2026-07-30' and m['record_count']==1
    r,_=parse_reference(page(),url=URL,retrieved_at='checked')
    assert r.pdf_url.iloc[0]==PDF and r.canonical_article_url.iloc[0]==FRIENDLY
    assert 'event_type' not in d and 'event_type' not in r

@pytest.mark.parametrize('url',[URL+'&p=456',URL+'&x=1',URL.replace('https','http'),URL.replace('moodyslocal.com.pe','other.test'),FRIENDLY+'?p=1'])
def test_invalid_action_urls_fail(url):
    with pytest.raises(InvalidQueryError):action_reference(url)

@pytest.mark.parametrize('transform',[lambda x:x.replace('</html>',''),lambda x:x.replace('30/07/2026','31/02/2026'),lambda x:x.replace('Título','Otro'),lambda x:x.replace('</tbody>',x[x.index('<tr><td>'):x.index('</tbody>')]+'</tbody>')])
def test_broken_index_fails(transform):
    with pytest.raises((SchemaChangedError,InvalidQueryError)):parse_index(transform(index()),retrieved_at='checked')

@pytest.mark.parametrize('transform',[lambda x:x.replace('?p=123','?p=456'),lambda x:x.replace('Jul 30','Feb 31'),lambda x:x.replace(PDF,PDF.replace('moodyslocal.com.pe','other.test')),lambda x:x.replace('</body>',f'<a href="{PDF}">Download</a></body>')])
def test_wrong_identity_date_or_ambiguous_download_fails(transform):
    with pytest.raises((SchemaChangedError,InvalidQueryError)):parse_reference(transform(page()),url=URL,retrieved_at='checked')

def test_friendly_action_cannot_redirect_to_unrelated_canonical():
    with pytest.raises(SchemaChangedError):parse_reference(page(),url=FRIENDLY.replace('ejemplo','otro'),retrieved_at='checked')

def test_index_only_offline_and_download_limit(tmp_path,monkeypatch):
    from fuentes_financieras.cli import retiros_clasificaciones as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=ActionIndexProvider(CATALOG['pe.moodys.indice_comunicados'])
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(url=INDEX,text=index()))
    assert p.sync().downloaded==1
    monkeypatch.setattr(cli,'source',lambda key:p if key=='pe.moodys.indice_comunicados' else pytest.fail('Unselected provider called'))
    args=['--descubrir','--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args+['--solo-indice'])==0
    book=load_workbook(tmp_path/'reports/retiros_clasificaciones.xlsx')
    assert book.sheetnames==['indice','seleccion','cobertura_indice']
    assert cli.main(args+['--desde','2027-01-01'])==0
    assert cli.main(args+['--desde','2027-01-01','--hasta','2026-01-01'])==1

def test_contradictory_reference_stops_before_pdf_provider(tmp_path,monkeypatch):
    from fuentes_financieras.cli import retiros_clasificaciones as cli
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    i=ActionIndexProvider(CATALOG['pe.moodys.indice_comunicados']);r=ActionReferencesProvider(CATALOG['pe.moodys.referencias_comunicados'])
    monkeypatch.setattr(i.transport,'request',lambda *a,**kw:SimpleNamespace(url=INDEX,text=index()))
    monkeypatch.setattr(r.transport,'request',lambda *a,**kw:SimpleNamespace(url=FRIENDLY,text=page().replace(TITLE,'Título diferente')))
    assert i.sync().downloaded==1 and r.sync(urls=[URL]).downloaded==1
    def source(key):
        if key=='pe.moodys.indice_comunicados':return i
        if key=='pe.moodys.referencias_comunicados':return r
        pytest.fail('PDF provider called with contradictory reference')
    monkeypatch.setattr(cli,'source',source)
    assert cli.main(['--descubrir','--load-only','--output-dir',str(tmp_path/'report')])==1
    assert not (tmp_path/'report/retiros_clasificaciones.xlsx').exists()
