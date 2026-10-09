"""Verify paginated inventory, coverage boundaries and explicit consumer selection."""
from types import SimpleNamespace
import pytest
import pandas as pd
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.providers.sbs.indice_noticias import parse_index, NewsIndexProvider, URL
from fuentes_financieras.providers.sbs.anuncios_regulatorios import announcement_reference
from fuentes_financieras.cli.anuncios_regulatorios import discovery_selection
from fuentes_financieras.exceptions import InvalidQueryError, SchemaChangedError


def html(page=1, ident='123', title='SBS interviene una entidad', month='JULIO'):
    return f'''<html><body><ul id="SBS_PressNote"><li class="list-news__item">
    <header><h3 class="list-news__item__header__title"><a title="{title}" href="/Noticia/DetalleNoticia?IdNoticia={ident}&amp;title=Titulo">Texto truncado...</a></h3>
    <div class="date">11 {month} 2024</div></header></li></ul>
    <ul class="pagination"><li class="active"><a href="http://www.sbs.gob.pe/Noticia/moduleId/99/pagina/{page}/controller/Item/action/Index">{page}</a></li>
    <li><a href="http://www.sbs.gob.pe/Noticia/moduleId/99/pagina/2/controller/%20/action/%20">Ultima</a></li></ul></body></html>'''


def parse(text, page=1):return parse_index(text, expected_page=page, source_url=URL, retrieved_at='checked')


def test_query_links_are_canonicalized_without_guessing_ids():
    identifier,url=announcement_reference('https://www.sbs.gob.pe/Noticia/DetalleNoticia?IdNoticia=123&title=Text')
    assert identifier=='123' and url.endswith('/idnoticia/123')
    for bad in ['IdNoticia=123&idnoticia=456','IdNoticia=123&IdNoticia=123','IdNoticia=123&other=x','title=only']:
        with pytest.raises(InvalidQueryError):announcement_reference('https://www.sbs.gob.pe/Noticia/DetalleNoticia?'+bad)


def test_full_title_date_and_page_are_source_fields_not_events():
    data,meta=parse(html())
    assert data.iloc[0].source_title=='SBS interviene una entidad'
    assert data.iloc[0].listed_date=='2024-07-11' and data.iloc[0].index_page==1
    assert meta['module_id']=='99' and meta['listed_pages_total']==2
    assert 'event_type' not in data


@pytest.mark.parametrize('text',[html(2),html().replace('class="active"',''),html(month='UNKNOWN'),html().replace('title="SBS interviene una entidad"',''),html().replace('</html>','')])
def test_wrong_page_missing_controls_or_incomplete_cards_fail(text):
    with pytest.raises(SchemaChangedError):parse(text)


def test_duplicate_or_external_cards_fail():
    text=html();card=text.split('<li class="list-news__item">')[1].split('</li>')[0]
    with pytest.raises(SchemaChangedError):parse(text.replace('</ul>', '<li class="list-news__item">'+card+'</li></ul>',1))
    with pytest.raises(SchemaChangedError):parse(text.replace('/Noticia/DetalleNoticia?', 'https://other.test/Noticia/DetalleNoticia?'))


def test_filter_is_inclusive_literal_and_does_not_infer_events():
    first,_=parse(html());second,_=parse(html(2,'124','Otra nota'),page=2)
    data=pd.concat([first,second])
    selected=discovery_selection(data,desde='2024-07-11',hasta='2024-07-11',palabras=['INTERVIENE'])
    assert selected.announcement_id.tolist()==['123']
    assert discovery_selection(data,palabras=['.*']).empty
    assert discovery_selection(data,desde='2025-01-01').empty
    with pytest.raises(ValueError):discovery_selection(data,desde='2025-01-01',hasta='2024-01-01')
    contradictory=first.copy();contradictory['source_title']='Changed'
    with pytest.raises(ValueError):discovery_selection(pd.concat([first,contradictory]))


def test_independent_provider_caches_pages_and_rejects_clamped_page(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=NewsIndexProvider(CATALOG['pe.sbs.indice_noticias']);calls=[]
    def request(*a,**kw):
        url=a[1];calls.append(url);page=2 if '/pagina/2/' in url else 1
        return SimpleNamespace(text=html(page,str(122+page)),url=url)
    monkeypatch.setattr(p.transport,'request',request)
    assert p.sync(paginas=[1,2,2]).downloaded==2 and len(calls)==2
    assert p.sync(paginas=[1,2]).skipped_existing==2 and len(calls)==2
    previous=p.load().copy()
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(text=html(),url=a[1]))
    assert p.sync(paginas=[2],force=True).failed==1 and p.load().equals(previous)
    with pytest.raises(InvalidQueryError):p.single_request(pagina=True)


def test_cli_empty_selection_exports_scope_without_article_requests(tmp_path,monkeypatch):
    from fuentes_financieras.cli import anuncios_regulatorios as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=NewsIndexProvider(CATALOG['pe.sbs.indice_noticias'])
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(text=html(),url=a[1]))
    assert p.sync().downloaded==1
    def source(dataset):
        assert dataset=='pe.sbs.indice_noticias'
        return p
    monkeypatch.setattr(cli,'source',source)
    def fail(*a,**kw):raise AssertionError('Unexpected network')
    monkeypatch.setattr(p.transport,'request',fail)
    args=['--descubrir','--load-only','--palabras','no-match','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args)==0
    book=load_workbook(tmp_path/'reports/anuncios_regulatorios.xlsx')
    assert book['seleccion'].max_row==1 and book['cobertura_indice'].max_row==2
    assert 'Alcance de la captura' in [c.value for c in book['cobertura_indice'][1]]
