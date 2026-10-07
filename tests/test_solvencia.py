"""Check mixed units, dates, missing totals and contradictory source values."""
from io import BytesIO
from datetime import datetime
from types import SimpleNamespace
import importlib.util
import sys

from openpyxl import Workbook
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.providers.sbs.solvencia import (
    SolvencyProvider, RATIO_METRICS, CAPITAL_METRICS, parse_ratios, parse_capital,
)


def workbook(kind='ratios', *, period='2026-07-31', unit=None, mismatch=False, missing=False, shifted=False):
    book = Workbook()
    sheet = book.active
    sheet.cell(2, 1, 'Requerimiento de Patrimonio Efectivo y Ratio de Capital Global' if kind == 'ratios' else 'Patrimonio Efectivo')
    sheet.cell(4, 1, datetime.fromisoformat(period))
    sheet.cell(5, 1, unit or '(En miles de soles)')
    sheet.cell(6, 1, datetime(2026, 9, 28))
    sheet.cell(9, 1, 'EMPRESAS' if kind == 'ratios' else 'ENTIDAD')
    metrics = RATIO_METRICS if kind == 'ratios' else CAPITAL_METRICS
    for col, (_, label) in metrics.items():
        row = 8 if kind == 'ratios' or col >= 3 else 9
        sheet.cell(row, col+1, label)
    if kind == 'ratios':
        sheet.cell(7, 2, 'REQUERIMIENTO DE PATRIMONIO EFECTIVO')
        sheet.cell(7, 6, 'ACTIVOS Y CONTINGENTES PONDERADOS')
        for col in (10, 11, 12):
            sheet.cell(10, col+1, '(En porcentaje)')
    else:
        # Fill the same 5-column layout as published workbooks.
        sheet.cell(8, 2, 'Patrimonio Efectivo de Nivel 1')
    for row, name in ((12, 'Entidad'), (13, 'TOTAL SISTEMA')):
        sheet.cell(row, 1, name)
        for col in metrics:
            value = {5: 100, 6: 20, 7: 30, 8: 150, 10: 12, 11: 13, 12: 17}.get(col, 1) if kind == 'ratios' else {1: 70, 2: 0, 3: 30, 4: 1000 if unit == '(En porcentaje)' else 100}[col]
            if mismatch and col == (8 if kind == 'ratios' else 1):
                value += 1
            if not (missing and col == (8 if kind == 'ratios' else 4)):
                sheet.cell(row, col+1, value)
    sheet.cell(15, 1, 'Fuente: Reportes SBS')
    if shifted:
        sheet.cell(8, 13, 'Otro ratio')
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def parse(content, kind='ratios'):
    return (parse_ratios if kind == 'ratios' else parse_capital)(content, entity_type='C', period='2026-07', source_url='test', retrieved_at='test')[0]


def test_ratio_units_and_effective_date():
    data = parse(workbook())
    assert len(data) == 20
    assert set(data.period_date) == {'2026-07-31'}
    assert set(data.source_auxiliary_date) == {'2026-09-28'}
    ratio = data[data.metric == 'global_capital_ratio']
    assert ratio.value.tolist() == [17, 17]
    assert set(ratio.unit) == {'percent'}
    amounts = data[data.metric == 'risk_weighted_assets_total']
    assert set(amounts.unit_multiplier) == {1000}
    assert set(data.entity_scope) == {'entity', 'system_aggregate'}


@pytest.mark.parametrize('kwargs', [{'period': '2026-06-30'}, {'unit': '(En millones de soles)'}, {'mismatch': True}, {'missing': True}, {'shifted': True}])
def test_rejects_changed_period_units_headers_or_apr_total(kwargs):
    with pytest.raises(SchemaChangedError):
        parse(workbook(**kwargs))


def test_capital_percentages_do_not_assign_total_currency_scale():
    data = parse(workbook('capital', unit='(En porcentaje)'), 'capital')
    totals = data[data.metric == 'effective_capital_total']
    assert set(totals.unit) == {'unspecified_by_source'}
    assert totals.unit_multiplier.isna().all()
    assert set(totals.data_quality_flags) == {'source_unit_unspecified'}
    assert set(data[data.metric != 'effective_capital_total'].unit) == {'percent'}


def test_preserves_published_component_mismatch():
    data = parse(workbook('capital', unit='(En porcentaje)', mismatch=True), 'capital')
    assert 'published_components_sum_mismatch' in data.data_quality_flags.iloc[0]
    assert data[data.metric == 'common_equity_tier1'].value.tolist() == [71, 71]


def test_provider_reuses_cache(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path))
    provider = SolvencyProvider(CATALOG['pe.sbs.solvencia'])
    url = 'https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Julio/C-1252-jl2026.XLS'
    calls = []
    def request(method, target):
        calls.append(target)
        return SimpleNamespace(text=f'<a href="{url}">Julio</a>', content=workbook())
    monkeypatch.setattr(provider.transport, 'request', request)
    assert provider.sync(desde='2026-07', tipos=['C']).downloaded == 1
    assert provider.sync(desde='2026-07', tipos=['C']).skipped_existing == 1
    assert len(calls) == 2
    assert provider.sync(desde='2026-08', tipos=['C']).unavailable == 1


def test_bootstrap_imports_package_created_in_new_user_site(tmp_path, monkeypatch):
    from pathlib import Path
    path = Path(__file__).resolve().parents[1] / 'scripts/_bootstrap.py'
    spec = importlib.util.spec_from_file_location('launcher_bootstrap', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    (tmp_path / 'pyproject.toml').write_text('[project]\ndependencies = ["fresh-demo-package>=1,<2"]\n')
    user_site = tmp_path / 'new_user_site'
    monkeypatch.setattr(module, 'ROOT', tmp_path)
    original_version = module.version
    def version(name):
        if name == 'fresh-demo-package':
            raise module.PackageNotFoundError(name)
        return original_version(name)
    monkeypatch.setattr(module, 'version', version)
    monkeypatch.setattr(module.site, 'ENABLE_USER_SITE', True)
    monkeypatch.setattr(module.site, 'getusersitepackages', lambda: str(user_site))
    def install(command):
        assert command[:3] == [sys.executable, '-m', 'pip']
        user_site.mkdir()
        (user_site / 'fresh_demo_package.py').write_text('VALUE = 42\n')
    monkeypatch.setattr(module.subprocess, 'check_call', install)
    old_path = list(sys.path)
    try:
        module.prepare()
        assert __import__('fresh_demo_package').VALUE == 42
    finally:
        sys.path[:] = old_path
        sys.modules.pop('fresh_demo_package', None)
