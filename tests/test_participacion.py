"""Guard published ranks, missing shares and comparison scope."""
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
import pytest
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.providers.sbs.participacion import parse_workbook,PRODUCTS


def workbook(*,dash=False,wrong_date=False,changed=False,foreign_note=True,mismatch=False):
    b=Workbook();s=b.active;s.append(['Ranking de Créditos, Depósitos y Patrimonio'])
    s.append([datetime(2026,7 if wrong_date else 8,31)]);s.append(['(En miles de soles)'])
    for k,label in enumerate(PRODUCTS):
        r=5+k*7;s.cell(r,1,label);s.cell(r+1,1,'Empresas')
        for c,v in [(4,'Monto'),(5,'Participación'),(6,'Porcentaje')]:s.cell(r+1,c,v)
        s.cell(r+2,5,'( % )');s.cell(r+2,6,'Acumulado')
        for off,(name,amount,share,total) in enumerate([('Entidad A',60,60,60),('Entidad B',40,40,100)],3):
            s.cell(r+off,1,off-2);s.cell(r+off,3,name);s.cell(r+off,4,amount)
            s.cell(r+off,5,'-' if dash and off==4 else share)
            s.cell(r+off,6,total+1 if mismatch else total)
    s.cell(27,1,'Fuente: Balance de Comprobación')
    if foreign_note:s.cell(28,1,'No incluye sucursales en el exterior')
    if changed:s.cell(6,4,'Importe nuevo')
    stream=BytesIO();b.save(stream);return stream.getvalue()


def parse(content):return parse_workbook(content,entity_type='B',period='2026-08',source_url='test',retrieved_at='test')[0]


def test_ranks_and_shares_are_published_not_recomputed():
    d=parse(workbook(dash=True));assert len(d)==18
    assert set(d.published_rank)=={1,2} and d.value.isna().sum()==3
    assert set(d[d.value.isna()].source_value_token)=={'-'}
    assert set(d.comparison_scope)=={'domestic_banking_excluding_foreign_branches'}
    assert set(d[d.metric=='amount'].unit)=={'thousands_PEN'}
    assert set(d[d.metric=='market_share'].unit)=={'percent'}


@pytest.mark.parametrize('kwargs',[{'wrong_date':True},{'changed':True},{'foreign_note':False}])
def test_ranking_guards_date_headers_and_territorial_scope(kwargs):
    with pytest.raises(SchemaChangedError):parse(workbook(**kwargs))


def test_cumulative_inconsistency_is_preserved_with_warning():
    d=parse(workbook(mismatch=True))
    assert 'published_cumulative_mismatch' in set(d.data_quality_flags)
    assert 101 in d.value.tolist()
