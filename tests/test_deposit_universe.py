import pandas as pd
import pytest
from openpyxl import load_workbook

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.entities import EntityCatalog, correspondence_report, observation_changes
from fuentes_financieras.exceptions import InvalidQueryError, SchemaChangedError
from fuentes_financieras.providers.sbs.deposit_universe import DepositUniverseProvider
from fuentes_financieras.cli import deposit_universe as cli


def capture():
    sections = ['Bancos', 'Financieras', 'Cajas municipales de ahorro y crédito',
                'Cajas rurales de ahorro y crédito']
    return ''.join('<div>' + title + '</div>' + ''.join(
        f'<span>- ENTIDAD {code} {i}</span>' for i in range(3))
        + '<p>Conoce a las empresas que cuentan con cobertura</p>'
        for code, title in zip('BFCR', sections)) + (
        '<div>Cooperativas de ahorro y crédito - Coopac</div>'
        '<span>COOPERATIVA DE AHORRO Y CREDITO PRUEBA</span>')


@pytest.fixture
def provider(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path / 'data'))
    return DepositUniverseProvider(CATALOG['pe.sbs.universo_depositos'])


def test_capture_cache_integrity_and_partial_page(provider, tmp_path):
    html = tmp_path / 'capture.html'
    html.write_text(capture())
    result = provider.import_capture(html, observed_on='2026-01-01')
    assert result.downloaded == 1
    frame = provider.load(latest=True)
    assert len(frame) == 12
    assert set(frame.entity_type_code) == set('BFCR')
    assert not frame.entity_name.str.contains('COOPERATIVA').any()
    assert provider.import_capture(html, observed_on='2026-01-01').skipped_existing == 1
    assert provider.import_capture(html, observed_on='2026-01-01', force=True).unchanged == 1
    with pytest.raises(SchemaChangedError):
        provider.from_html('<h1>Access denied</h1>', observed_on='2026-01-01')
    assert len(provider.load()) == 12
    with pytest.raises(InvalidQueryError):
        provider.fetch(fecha='2026-01-01')
    assert provider.load(desde='2026-01-02').empty


def test_identities_dates_conflicts_and_changes(provider, tmp_path):
    first = provider.from_html(capture(), observed_on='2026-01-01').data
    catalog = EntityCatalog(tmp_path / 'identities.json')
    catalog.observe(first)
    identity = catalog.records[0]['entity_id']
    catalog.save()
    reopened = EntityCatalog(catalog.path)
    assert reopened.records[0]['entity_id'] == identity
    assert reopened.resolve(dataset='rates', type_code='B', name='ENTIDAD B 0',
                            period='2026-01-01')[0] == identity
    assert reopened.resolve(dataset='rates', type_code='F', name='ENTIDAD B 0',
                            period='2026-01-01')[1] == 'unmatched'
    alias = dict(dataset='pe.sbs.tasas_pasivas', entity_type_code='B', alias='OTRO NOMBRE',
                 entity_id=identity, valid_from='2026-02-01', valid_to=None, evidence='documento')
    rates = pd.DataFrame([dict(entity_type='B', entity_name='OTRO NOMBRE', period_date='2026-01-01'),
                          dict(entity_type='B', entity_name='OTRO NOMBRE', period_date='2026-02-01')])
    report = correspondence_report(reopened, rates=rates, aliases=[alias])
    assert report.match_status.tolist() == ['unmatched', 'alias']
    ratings = pd.DataFrame([dict(entity_type_code='B', entity_name='ENTIDAD B 0',
                                period_date='2026-03-01')])
    assert correspondence_report(reopened, ratings=ratings).entity_id.iloc[0] == identity
    conflict = dict(alias, entity_id=catalog.records[1]['entity_id'])
    assert correspondence_report(reopened, rates=rates, aliases=[alias, conflict]).match_status.iloc[1] == 'ambiguous'
    second = first.iloc[1:].copy()
    second['period'] = '2026-02-01'
    changes = observation_changes(pd.concat([first, second]), reopened)
    assert changes.change.tolist() == ['disappeared']


def test_explicit_rename_preserves_identity(provider, tmp_path):
    first = provider.from_html(capture(), observed_on='2026-01-01').data
    catalog = EntityCatalog(tmp_path / 'identities.json')
    catalog.observe(first)
    identity = catalog.records[0]['entity_id']
    second = first.copy()
    second['period'] = '2026-02-01'
    second.loc[0, 'entity_name'] = 'NUEVO NOMBRE'
    alias = dict(dataset='pe.sbs.universo_depositos', entity_type_code='B', alias='NUEVO NOMBRE',
                 entity_id=identity, valid_from='2026-02-01', valid_to=None, evidence='documento')
    catalog.observe(second, aliases=[alias])
    assert len(catalog.records) == 12
    assert observation_changes(pd.concat([first, second]), catalog, [alias]).empty


def test_cli_reports_missing_local_datasets(provider, tmp_path):
    html = tmp_path / 'capture.html'
    html.write_text(capture())
    out = tmp_path / 'reports'
    assert cli.run(['--html', str(html), '--observed-on', '2026-01-01',
                    '--output-dir', str(out)]) == 0
    wb = load_workbook(out / 'universo_correspondencias.xlsx')
    assert wb['Universo vigente observado'].freeze_panes == 'A2'
    assert list(wb['Universo vigente observado'].tables.values())[0].tableStyleInfo.name == 'TableStyleLight9'
    assert wb['Cobertura local']['C2'].value == 'Sin datos locales; no evaluado'
    assert wb['Correspondencias'].max_row == 1


def test_reviewed_aliases_aggregates_scope_and_date_bounds(tmp_path):
    universe = pd.DataFrame([
        dict(entity_type_code='B', entity_name='BANBIF', period='2026-10-07', source_url='https://www.sbs.gob.pe'),
        dict(entity_type_code='B', entity_name='SANTANDER PERU', period='2026-10-07', source_url='https://www.sbs.gob.pe'),
        dict(entity_type_code='B', entity_name='BN. SANTANDER CONS.', period='2026-10-07', source_url='https://www.sbs.gob.pe'),
    ])
    catalog = EntityCatalog(tmp_path / 'identities.json')
    catalog.observe(universe)
    rates = pd.DataFrame([
        dict(entity_type='B', entity_name=name, period_date='2026-10-06')
        for name in ['BIF', 'Santander', 'Santander Cons. Bank', 'Promedio', 'Mitsui']
    ])
    report = correspondence_report(catalog, rates=rates)
    assert report.match_status.tolist() == ['alias', 'alias', 'alias', 'aggregate', 'unmatched']
    assert report.entity_id.iloc[1] != report.entity_id.iloc[2]
    assert report.evidence.iloc[:3].str.contains('https://www.sbs.gob.pe').all()
    rates['period_date'] = '2020-01-01'
    assert correspondence_report(catalog, rates=rates).match_status.iloc[0] == 'unmatched'
    ratings = pd.DataFrame([dict(entity_type_code='S', entity_name='INSURANCE', period_date='2026-09-30')])
    assert correspondence_report(catalog, ratings=ratings).match_status.iloc[0] == 'outside_scope'
