"""Guard annualization, denominators, monetary productivity units and raw names."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.models import SyncResult
from fuentes_financieras.providers.sbs.rentabilidad import (
    ANNUAL_NOTE, TITLES, PROFIT_METRICS, EFFICIENCY_METRICS, observation_window,
    parse_profitability, parse_efficiency, ProfitabilityProvider,
)


def indicators(*, tipo='B', date=datetime(2026,8,31), dash=False, changed=False,
               annual_note=ANNUAL_NOTE, auxiliary=False, percentage_format=False, repeated=True,
               unnamed=False, stray=False, duplicate=False):
    b=Workbook();s=b.active
    for col in ((1,4) if repeated else (1,)):
        s.cell(1,col,TITLES[tipo]);s.cell(2,col,date)
        if tipo in 'BF':s.cell(3,col,'(En porcentaje)')
        s.cell(5,col+1,'Entidad' if col==1 else 'TOTAL SISTEMA')
        if auxiliary:s.cell(6,col+1,'Nombre auxiliar sin fecha')
        s.cell(8,col,'SOLVENCIA')
        s.cell(10,col,'EFICIENCIA Y GESTIÓN')
        for row,label in enumerate(EFFICIENCY_METRICS[tipo],11):
            s.cell(row,col,label+(' (%)' if tipo in 'CR' and '(Miles' not in label else ''))
            value=125 if row==12 else 1000 if 'Miles' in label else 5
            s.cell(row,col+1,value)
        s.cell(18,col,'RENTABILIDAD')
        for row,label in enumerate(PROFIT_METRICS[tipo],19):
            s.cell(row,col,label+(' (%)' if tipo in 'CR' else ''))
            s.cell(row,col+1,'-' if dash and row==19 else -2)
            if percentage_format:s.cell(row,col+1).number_format='0%'
        s.cell(22,col,'LIQUIDEZ')
        s.cell(25,col,'Nota: Definiciones del glosario')
        if annual_note:s.cell(26,col,annual_note)
    if changed:s.cell(11,1,'Indicador nuevo no revisado')
    if unnamed:s['B5']=None
    if stray:s.cell(17,2,999)
    if duplicate:s.cell(21,1,next(iter(PROFIT_METRICS[tipo])))
    stream=BytesIO();b.save(stream);return stream.getvalue()


def parse(content, parser=parse_profitability, tipo='B'):
    return parser(content,entity_type=tipo,period='2026-08',source_url='test',retrieved_at='test')[0]


@pytest.mark.parametrize('tipo',list('BFCR'))
def test_profitability_is_twelve_months_and_keeps_negative_values_and_dash(tipo):
    d=parse(indicators(tipo=tipo,dash=True),tipo=tipo)
    assert len(d)==4 and set(d.metric)=={'return_on_equity','return_on_assets'}
    assert d.value.isna().sum()==2 and set(d.value.dropna())=={-2}
    assert set(d.observation_start)=={'2025-09-01'} and set(d.observation_end)=={'2026-08-31'}
    assert set(d.measurement_basis)=={'rolling_12_months'} and set(d.unit)=={'percent'}
    assert set(d.annualization_definition)=={ANNUAL_NOTE}
    assert set(d[d.value.isna()].source_value_token)=={'-'}
    expected={'annualized_profit_as_labeled'} if tipo=='F' else {'annualized_net_profit'}
    assert set(d.numerator_basis)==expected


def test_efficiency_does_not_harmonize_denominators_or_infer_unqualified_periods():
    b=parse(indicators(),parse_efficiency)
    c=parse(indicators(tipo='C'),parse_efficiency,'C')
    assert len(b)==12 and len(c)==12
    admin_b=b[b.metric=='administrative_expenses_productive_assets_ratio']
    admin_c=c[c.metric=='administrative_expenses_average_credit_ratio']
    assert set(admin_b.denominator_basis)=={'average_productive_assets'}
    assert set(admin_c.denominator_basis)=={'average_direct_and_indirect_credit'}
    unqualified=b[b.metric=='operating_expenses_financial_margin_ratio']
    assert set(unqualified.measurement_basis)=={'unspecified_by_source'}
    assert set(unqualified.observation_start)=={''} and set(unqualified.observation_end)=={''}
    assert set(unqualified.value)=={125} and set(unqualified.data_quality_flags)=={''}
    annualized=c[c.metric=='annualized_operating_expenses_financial_margin_ratio']
    assert set(annualized.measurement_basis)=={'rolling_12_months'}
    assert set(annualized.denominator_basis)=={'annualized_total_financial_margin'}
    productivity=b[b.metric=='direct_credit_per_person']
    assert set(productivity.unit)=={'thousands_PEN_per_person'} and set(productivity.unit_multiplier)=={1000}
    assert set(productivity.observation_start)=={'2026-08-31'}
    assert set(c[c.metric=='direct_credit_per_office'].unit)=={'thousands_PEN_per_office'}


def test_auxiliary_name_is_not_used_as_entity_identity():
    d=parse(indicators(auxiliary=True))
    assert set(d.entity_name)=={'Entidad','TOTAL SISTEMA'}
    assert set(d.source_auxiliary_entity_name)=={'Nombre auxiliar sin fecha'}
    assert set(d.data_quality_flags)=={'source_auxiliary_entity_name'}


@pytest.mark.parametrize('kwargs',[{'annual_note':''},{'annual_note':'Los valores anualizados se multiplican por doce.'},
    {'date':datetime(2026,7,31)},{'percentage_format':True},{'unnamed':True},{'duplicate':True}])
def test_profitability_rejects_ambiguous_dates_units_methodology_or_entities(kwargs):
    with pytest.raises(SchemaChangedError):parse(indicators(**kwargs))


@pytest.mark.parametrize('kwargs',[{'changed':True},{'stray':True}])
def test_efficiency_rejects_unknown_metrics_and_values_without_metric(kwargs):
    with pytest.raises(SchemaChangedError):parse(indicators(**kwargs),parse_efficiency)


def test_book_entity_type_must_match_query():
    with pytest.raises(SchemaChangedError):parse(indicators(tipo='C'),tipo='B')


def test_annual_window_at_year_boundary():
    assert observation_window('2025-12','rolling_12_months')==('2025-01-01','2025-12-31')
    assert observation_window('2026-01','rolling_12_months')==('2025-02-01','2026-01-31')


def test_profitability_cache_skips_network_and_preserves_valid_capture_on_failure(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=ProfitabilityProvider(CATALOG['pe.sbs.rentabilidad']);calls=[]
    url='https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-2401-ag2026.XLS'
    def request(method,target):
        calls.append(target);return SimpleNamespace(text=f'<a href="{url}">Agosto</a>',content=indicators())
    monkeypatch.setattr(p.transport,'request',request)
    assert p.sync(desde='2026-08',tipos=['B']).downloaded==1
    assert p.sync(desde='2026-08',tipos=['B']).skipped_existing==1 and len(calls)==2
    previous=p.load().copy()
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(text=f'<a href="{url}">Agosto</a>',content=indicators(annual_note='')))
    assert p.sync(desde='2026-08',tipos=['B'],force=True).failed==1
    assert p.load().equals(previous)
    assert p.sync(desde='2026-09',tipos=['B']).unavailable==1


def test_cli_reports_in_spanish_and_rejects_incomplete_selection(tmp_path,monkeypatch):
    from fuentes_financieras.cli import fondeo, rentabilidad
    d=parse(indicators())
    req=SimpleNamespace(params={'tipo':'B','periodo':'2026-08'})
    fake=SimpleNamespace(plan_sync=lambda **kw:[req],sync=lambda **kw:SyncResult('test',requested=1,downloaded=1),load=lambda **kw:d)
    monkeypatch.setattr(fondeo,'source',lambda dataset:fake)
    args=['--datasets','rentabilidad','--desde','2026-08','--tipos','B','--output-dir',str(tmp_path)]
    assert rentabilidad.main(args)==0
    path=tmp_path/'rentabilidad_eficiencia.xlsx';b=load_workbook(path);s=b.active
    assert s.freeze_panes=='A2' and next(iter(s.tables.values())).tableStyleInfo.name=='TableStyleLight9'
    assert not s.column_dimensions
    assert 'Base del denominador' in [c.value for c in s[1]]
    path.unlink()
    fake.sync=lambda **kw:SyncResult('test',requested=1,unavailable=1)
    assert rentabilidad.main(args)==1 and not path.exists()
    fake.load=lambda **kw:d.iloc[:0]
    assert rentabilidad.main(args+['--load-only'])==1
