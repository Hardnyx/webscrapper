"""Verify report provenance, strict summary structure and WebForms framing."""
from types import SimpleNamespace
import pandas as pd
import pytest
from openpyxl import load_workbook
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError, SourceUnavailableError
from fuentes_financieras.models import SyncResult
from fuentes_financieras.storage import content_hash, schema_hash
from fuentes_financieras.providers.sbs._clasificaciones_parser import to_long_form
from fuentes_financieras.providers.sbs.clasificaciones_riesgo import RiskRatingsProvider, _parse_delta_response
from fuentes_financieras.providers.sbs.informes_riesgo import report_inventory

URL='https://www.sbs.gob.pe/app/iece/paginas/MostrarResumenClasificaciones.aspx'
LINK='https://extranet.sbs.gob.pe/iece/descargar?codClasificadora=001196&codPeriodo=202601&numArchivo=22&numVersion=1'
QUERY=dict(period_code='202601',type_code_by_label={'Banco':'B'},source_url=URL,retrieved_at='test')


def html(*, link=LINK, icon='', duplicate=False, short=False, multiple=False, plain=False):
    cells=f'<td><a href="{link}">A-</a>{icon}</td>'
    if plain:cells='<td>A-</td>'
    if multiple:cells='<td><a>A-</a><a>B</a></td>'
    row='<tr><td>Banco</td><td>Entidad</td>'+('' if short else cells)+'</tr>'
    return '<table><tr><th>Tipo de Entidad</th><th>Entidad</th><th>Clasificadora</th></tr>'+row+(row if duplicate else '')+'<tr class="spacer"><td colspan="3"></td></tr></table>'


def test_report_identity_preserves_codes_versions_and_does_not_infer_deposit_rating_or_date():
    d=to_long_form(html(),**QUERY)
    assert d.iloc[0].report_id=='001196:202601:22:1' and d.iloc[0].report_url==LINK
    assert d.iloc[0].rating_kind=='institutional_summary' and pd.isna(d.iloc[0].trend)
    inventory=report_inventory(html(),**QUERY)
    assert inventory.iloc[0].summary_rating=='A-' and inventory.iloc[0].report_date==''
    assert inventory.iloc[0].document_status=='linked_not_downloaded' and 'rating' not in inventory
    changed=report_inventory(html(link=LINK.replace('numVersion=1','numVersion=2')),**QUERY)
    assert changed.iloc[0].report_id.endswith(':2')


def test_known_change_is_not_outlook_and_unknown_symbol_is_preserved():
    known='<img src="../Imagenes/SUBIO.png" title="Subió en relación a la clasificación anterior">'
    d=to_long_form(html(icon=known),**QUERY)
    assert d.iloc[0].trend=='up' and d.iloc[0].trend_basis=='published_change_vs_previous_classification'
    assert d.iloc[0].source_change_title=='Subió en relación a la clasificación anterior'
    unknown=to_long_form(html(icon='<img src="nuevo.png" alt="Estable">'),**QUERY)
    assert pd.isna(unknown.iloc[0].trend) and unknown.iloc[0].source_change_alt=='Estable'
    assert unknown.iloc[0].data_quality_flags=='source_change_symbol_unrecognized'
    absent=to_long_form(html(link=''),**QUERY)
    assert absent.iloc[0].data_quality_flags=='source_report_link_missing'


@pytest.mark.parametrize('options',[
    {'duplicate':True},{'short':True},{'multiple':True},{'plain':True},
    {'link':LINK.replace('202601','202602')},{'link':LINK+'&numArchivo=99'},
    {'link':'javascript:alert(1)'},{'link':LINK.replace('extranet.sbs.gob.pe','example.org')},
    {'icon':'<img src="SUBIO.png" title="Bajó en relación a la clasificación anterior">'},
])
def test_summary_rejects_ambiguous_or_wrong_period_records(options):
    with pytest.raises(SchemaChangedError):to_long_form(html(**options),**QUERY)


def segment(kind,identifier,content):
    size=len(content.encode('utf-16-le'))//2
    return f'{size}|{kind}|{identifier}|{content}|'


def test_delta_lengths_preserve_pipes_crlf_hidden_values_and_non_bmp_characters():
    content='<table>literal |10|hiddenField|__VIEWSTATE|fake|\r\n😀</table>'
    raw=segment('updatePanel','ctl00_MainContent_UpTblResumen',content)+segment('hiddenField','__VIEWSTATE','a|b')
    fragment,hidden=_parse_delta_response(raw)
    assert fragment==content and hidden=={'__VIEWSTATE':'a|b'}


@pytest.mark.parametrize('raw',[
    '20|updatePanel|ctl00_MainContent_UpTblResumen|short|',
    segment('updatePanel','other','x'),segment('error','500','error'),
    segment('updatePanel','ctl00_MainContent_UpTblResumen','x')*2,
    segment('updatePanel','ctl00_MainContent_UpTblResumen','x')+segment('hiddenField','__VIEWSTATE','a')*2,
])
def test_delta_rejects_truncation_errors_duplicate_panels_and_state(raw):
    with pytest.raises(SourceUnavailableError):_parse_delta_response(raw)


def test_contract_migrates_and_failed_refresh_preserves_previous_capture(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=RiskRatingsProvider(CATALOG['pe.sbs.clasificaciones_riesgo'])
    monkeypatch.setattr(p,'available_periods',lambda:[{'period_code':'202601','period_date':'2026-03-31','year':2026,'semester':1}])
    p.client.type_code_by_label={'Banco':'B'};p.client.current_period='202601'
    calls=[]
    def fetch(code):calls.append(code);return html(),html()
    monkeypatch.setattr(p.client,'fetch_period',fetch)
    assert p.sync(periodos=['202601']).downloaded==1
    entry=p.storage.manifest.get('202601')
    legacy=p.load().drop(columns=['rating_kind','trend_basis','source_change_icon_url','source_change_title','source_change_alt',
        'data_quality_flags','report_url','report_agency_code','report_period_code','report_file_number','report_version','report_id'])
    p.storage.upsert_period('year=2026','202601',legacy)
    entry.update(contract_version='2',content_hash=content_hash(legacy),schema_hash=schema_hash(legacy))
    p.storage.manifest.last_schema_hash=schema_hash(legacy)
    p.storage.manifest.last_schema_contract_version='2';p.storage.manifest.save()
    assert p.sync(periodos=['202601']).downloaded==1 and len(calls)==2
    assert p.storage.manifest.get('202601')['contract_version']=='3'
    assert p.sync(periodos=['202601']).skipped_existing==1 and len(calls)==2
    previous=p.load().copy()
    monkeypatch.setattr(p.client,'fetch_period',lambda code:(html(duplicate=True),html(duplicate=True)))
    assert p.sync(periodos=['202601'],force=True).failed==1 and p.load().equals(previous)


def test_cli_selected_periods_excel_and_incomplete_capture(tmp_path,monkeypatch):
    from fuentes_financieras.cli import clasificaciones_riesgo as cli
    data=to_long_form(html(),**QUERY)
    fake=SimpleNamespace(plan_sync=lambda **kw:[SimpleNamespace(period_key='202601')],
        sync=lambda **kw:SyncResult('test',requested=1,downloaded=1),load=lambda **kw:data)
    monkeypatch.setattr(cli,'source',lambda dataset:fake)
    args=['--periodos','202601','--no-second-sync','--output-dir',str(tmp_path)]
    assert cli.main(args)==0
    path=tmp_path/'clasificaciones_informes.xlsx';s=load_workbook(path).active
    assert s.freeze_panes=='A2' and not s.column_dimensions
    assert next(iter(s.tables.values())).tableStyleInfo.name=='TableStyleLight9'
    assert 'URL informe' in [c.value for c in s[1]]
    path.unlink();fake.sync=lambda **kw:SyncResult('test',requested=1,failed=1)
    assert cli.main(args)==1 and not path.exists()


def test_load_only_rejects_legacy_contract_without_network(tmp_path,monkeypatch):
    from fuentes_financieras.cli import clasificaciones_riesgo as cli
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=RiskRatingsProvider(CATALOG['pe.sbs.clasificaciones_riesgo'])
    monkeypatch.setattr(p,'available_periods',lambda:[{'period_code':'202601','period_date':'2026-03-31','year':2026,'semester':1}])
    p.client.type_code_by_label={'Banco':'B'}
    monkeypatch.setattr(p.client,'fetch_period',lambda code:(html(),html()))
    assert p.sync(periodos=['202601']).downloaded==1
    monkeypatch.setattr(cli,'source',lambda dataset:p)
    def fail(*a,**kw):raise AssertionError('Unexpected offline request')
    monkeypatch.setattr(p.client.transport,'request',fail)
    args=['--periodos','202601','--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args)==0
    (tmp_path/'reports/clasificaciones_informes.xlsx').unlink()
    p.storage.manifest.get('202601')['contract_version']='2'
    assert cli.main(args)==1 and not (tmp_path/'reports/clasificaciones_informes.xlsx').exists()
