"""Verify conservative cover extraction and atomic, integrity-checked PDF caching."""
from io import BytesIO
from types import SimpleNamespace
import pytest
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import InvalidQueryError, SourceUnavailableError, SchemaChangedError
from fuentes_financieras.providers.sbs.documentos_riesgo import (
    RiskDocumentsProvider, extract_cover_fields, parse_document, normalized_date,
)
from fuentes_financieras.providers.sbs._clasificaciones_parser import report_reference

URL='https://extranet.sbs.gob.pe/iece/descargar?codClasificadora=000409&codPeriodo=202601&numArchivo=30&numVersion=1'
PCR='''Fecha de Comité: 30 de marzo de 2026
Fortaleza Financiera
PEB+
Depósitos de Corto Plazo
PECategoría II
Depósitos de Mediano y Largo Plazo
PEA
Significado De La Calificación
Historial de Calificaciones
30/06/25 19/09/25 Fortaleza Financiera PEB Estable
Perspectiva
Estable.
'''


def pdf(text=PCR,encrypted=False):
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),
        NameObject('/BaseFont'):NameObject('/Helvetica'),NameObject('/Encoding'):NameObject('/WinAnsiEncoding')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();commands=['BT /F1 11 Tf 40 740 Td 15 TL']
    for line in text.splitlines():
        escaped=line.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
        commands.append(f'({escaped}) Tj T*')
    stream.set_data(('\n'.join(commands)+'\nET').encode('cp1252'))
    page[NameObject('/Contents')]=writer._add_object(stream)
    if encrypted:writer.encrypt('test-password')
    out=BytesIO();writer.write(out);return out.getvalue()


def reference():return report_reference(URL,'202601',URL)[0]


def test_pcr_cards_do_not_use_historical_rating_or_change_scale():
    fields=extract_cover_fields(PCR,agency_code='000409')
    assert len(fields)==5
    short=next(f for f in fields if f['field_kind']=='short_term_deposits')
    assert short['value_raw']=='PECategoría II' and short['temporal_role']=='current'
    outlook=next(f for f in fields if f['field_kind']=='outlook')
    assert outlook['value_raw']=='Estable.' and outlook['normalized_value']=='Estable'
    assert next(f for f in fields if f['field_kind']=='committee_date')['normalized_value']=='2026-03-30'


def test_current_previous_header_and_ambiguous_committee_dates():
    text='''Ratings Actual Anterior
Fortaleza Financiera1 B+ B
Con información auditada
1Clasificaciones otorgadas en Comités de fecha
18/03/2026 y 17/09/2025
Perspectiva
Positiva'''
    fields=extract_cover_fields(text,agency_code='000408')
    assert [f['value_raw'] for f in fields[:2]]==['B+','B']
    assert [f['temporal_role'] for f in fields[:2]]==['current','previous']
    dates=[f for f in fields if f['field_kind']=='committee_date']
    assert len(dates)==2 and all(f['extraction_status']=='needs_review' and f['temporal_role']=='unspecified' for f in dates)


def test_jcr_multiline_table_and_explicit_current_committee():
    text='''Rating Actual* Anterior**
Fortaleza
Financiera A- B+
Depósitos de Corto
Plazo CP1- CP2+
Depósitos de
Mediano y Largo
Plazo A A
*Información auditada a diciembre de 2025.
Aprobado en comité de 14-04-2026.
**Información a junio de 2025.
Perspectiva Estable Negativa'''
    fields=extract_cover_fields(text,agency_code='001196')
    assert len(fields)==9
    assert next(f for f in fields if f['field_kind']=='committee_date')['normalized_value']=='2026-04-14'


def test_moodys_keeps_entity_emitter_and_deposit_scale_separate():
    text='''Fecha de comité: 16 de marzo de 2026
Fecha de publicación: 27 de marzo de 2026
CLASIFICACIONES ACTUALES (*)
Clasificación Perspectiva
Entidad A Estable
Emisor AA.pe Estable
Depósitos de Corto Plazo ML A-1.pe -
(*) Nota de escala local.'''
    fields=extract_cover_fields(text,agency_code='000406')
    assert len(fields)==7
    assert next(f for f in fields if f['field_kind']=='short_term_deposits')['value_raw']=='ML A-1.pe'
    assert not any(f['field_kind']=='short_term_deposits_outlook' for f in fields)


def test_no_narrative_rating_is_promoted_to_structured_value():
    assert not extract_cover_fields('Depósitos de Corto Plazo antes CP2 y ahora CP1; perspectiva estable.',agency_code='000409')
    assert not extract_cover_fields(PCR,agency_code='999999')
    duplicate=PCR+'\nPerspectiva\nNegativa'
    assert all(f['extraction_status']=='needs_review' for f in extract_cover_fields(duplicate,agency_code='000409') if f['field_kind']=='outlook')


def test_pdf_has_page_evidence_digest_and_explicit_ocr_state():
    d,m=parse_document(pdf(),reference=reference(),retrieved_at='test')
    assert len(d)==5 and set(d.page_number)=={1} and d.evidence_text.ne('').all()
    assert m['page_count']==1 and set(d.pdf_sha256)=={m['pdf_sha256']}
    d,m=parse_document(pdf(''),reference=reference(),retrieved_at='test')
    assert d.iloc[0].extraction_status=='needs_ocr' and d.iloc[0].value_raw==''
    d,m=parse_document(pdf('Texto sin plantilla conocida'),reference=reference(),retrieved_at='test')
    assert d.iloc[0].extraction_status=='unsupported_cover'


@pytest.mark.parametrize('content',[b'<html>Access denied</html>',b'%PDF-1.7 truncated',pdf()[:-20]])
def test_incomplete_or_html_response_is_not_saved_as_pdf(content):
    with pytest.raises((SourceUnavailableError,SchemaChangedError)):parse_document(content,reference=reference(),retrieved_at='test')


def test_encrypted_pdf_is_rejected():
    with pytest.raises(SchemaChangedError):parse_document(pdf(encrypted=True),reference=reference(),retrieved_at='test')


@pytest.mark.parametrize('raw',['31/02/2026','99 de marzo de 2026','2026-03','septiembre 2026'])
def test_date_requires_complete_valid_calendar_evidence(raw):assert normalized_date(raw)==''


def test_provider_cache_pdf_integrity_and_failure_preserve_previous_file(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=RiskDocumentsProvider(CATALOG['pe.sbs.documentos_riesgo']);calls=[];content=pdf()
    def request(*a,**kw):calls.append(a);return SimpleNamespace(content=content,url=URL)
    monkeypatch.setattr(p.transport,'request',request)
    assert p.sync(urls=[URL,URL]).downloaded==1 and len(calls)==1
    assert p.sync(urls=[URL]).skipped_existing==1 and len(calls)==1
    path=p.pdf_path(reference()['report_id']);path.write_bytes(b'corrupt')
    assert p.sync(urls=[URL]).unchanged==1 and len(calls)==2 and path.read_bytes()==content
    previous=p.load().copy()
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(content=b'<html>error</html>',url=URL))
    assert p.sync(urls=[URL],redownload=True).failed==1
    assert p.load().equals(previous) and path.read_bytes()==content
    assert not path.with_suffix('.pdf.tmp').exists()


def test_explicit_version_change_and_official_url_validation(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=RiskDocumentsProvider(CATALOG['pe.sbs.documentos_riesgo'])
    requests=list(p.plan_sync(urls=[URL,URL.replace('numVersion=1','numVersion=2')]))
    assert len(requests)==2 and requests[0].period_key!=requests[1].period_key
    with pytest.raises(InvalidQueryError):p.single_request(url=URL.replace('extranet.sbs.gob.pe','other.test'))
    with pytest.raises(InvalidQueryError):list(p.plan_sync())


def test_cli_offline_evidence_report_and_corrupt_pdf_rejection(tmp_path,monkeypatch):
    from fuentes_financieras.cli import documentos_riesgo as cli
    from fuentes_financieras.providers.sbs.informes_riesgo import RiskReportInventoryProvider
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    inventory=RiskReportInventoryProvider(CATALOG['pe.sbs.informes_riesgo'])
    monkeypatch.setattr(inventory,'available_periods',lambda:[{'period_code':'202601','period_date':'2026-03-31','year':2026,'semester':1}])
    inventory.client.type_code_by_label={'Banco':'B'}
    html=f'<table><tr><th>Tipo de Entidad</th><th>Entidad</th><th>Agencia</th></tr><tr><td>Banco</td><td>Entidad</td><td><a href="{URL}">B+</a></td></tr></table>'
    monkeypatch.setattr(inventory.client,'fetch_period',lambda code:(html,html))
    assert inventory.sync(periodos=['202601']).downloaded==1
    documents=RiskDocumentsProvider(CATALOG['pe.sbs.documentos_riesgo'])
    monkeypatch.setattr(documents.transport,'request',lambda *a,**kw:SimpleNamespace(content=pdf(),url=URL))
    assert documents.sync(urls=[URL]).downloaded==1
    monkeypatch.setattr(cli,'source',lambda dataset:inventory if dataset=='pe.sbs.informes_riesgo' else documents)
    def fail(*a,**kw):raise AssertionError('Unexpected offline request')
    monkeypatch.setattr(documents.transport,'request',fail);monkeypatch.setattr(inventory.client.transport,'request',fail)
    args=['--periodos','202601','--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args)==0
    path=tmp_path/'reports/documentos_riesgo.xlsx';b=load_workbook(path)
    for s in b:
        assert s.freeze_panes=='A2' and not s.column_dimensions
        assert next(iter(s.tables.values())).tableStyleInfo.name=='TableStyleLight9'
    assert 'Texto de evidencia' in [c.value for c in b['campos'][1]]
    path.unlink();documents.pdf_path(reference()['report_id']).write_bytes(b'corrupt')
    assert cli.main(args)==1 and not path.exists()

    monkeypatch.setattr(documents.transport,'request',lambda *a,**kw:SimpleNamespace(content=pdf(),url=URL))
    assert cli.main([a for a in args if a!='--load-only']+['--redownload'])==0
    assert path.exists()


AAI_DEPOSITS = '''Ratings Actual Anterior
Fortaleza Financiera A+ A
Depósitos CP CP-1+ (pe) CP-1+ (pe)
Depósitos LP AAA (pe) AAA (pe)
Certificado de Depósitos Negociables
Segundo Programa CP-1+ (pe) CP-1+ (pe)
Con información financiera auditada a diciembre 2025.
Clasificaciones otorgadas en Comités de fechas 26/03/2026 y
29/09/2025.
Perspectiva
Estable'''


def test_aai_local_deposit_scales_remain_distinct_from_certificates_and_outlook():
    fields = extract_cover_fields(AAI_DEPOSITS, agency_code='000408')
    deposits = [f for f in fields if f['field_kind'] in ('short_term_deposits','long_term_deposits')]
    assert [(f['field_kind'],f['value_raw'],f['temporal_role']) for f in deposits] == [
        ('short_term_deposits','CP-1+ (pe)','current'), ('short_term_deposits','CP-1+ (pe)','previous'),
        ('long_term_deposits','AAA (pe)','current'), ('long_term_deposits','AAA (pe)','previous')]
    assert all(f['normalized_value'] == f['value_raw'] for f in deposits)
    assert len([f for f in fields if f['field_kind']=='outlook']) == 1
    assert not any(f['field_kind'].endswith('deposits_outlook') for f in fields)
    dates = [f for f in fields if f['field_kind']=='committee_date']
    assert {f['normalized_value'] for f in dates} == {'2026-03-26','2025-09-29'}
    assert all(f['extraction_status']=='needs_review' and f['temporal_role']=='unspecified' for f in dates)


@pytest.mark.parametrize('text', [
    AAI_DEPOSITS.replace('Depósitos CP','Certificados CP').replace('Depósitos LP','Bonos LP'),
    AAI_DEPOSITS.replace('(pe)','(us)'),
    AAI_DEPOSITS.replace('CP-1+ (pe) CP-1+ (pe)','CP-1+ (pe)').replace('AAA (pe) AAA (pe)','AAA (pe)'),
    AAI_DEPOSITS.replace('Ratings Actual Anterior','Historial de calificaciones'),
    AAI_DEPOSITS.replace('Ratings Actual Anterior','Ratings Actual Anterior\nRatings Actual Anterior'),
])
def test_aai_does_not_promote_wrong_instruments_scales_or_incomplete_pairs(text):
    fields = extract_cover_fields(text, agency_code='000408')
    assert not any(f['field_kind'] in ('short_term_deposits','long_term_deposits') for f in fields)


def test_jcr_alternative_labels_and_inline_current_committee():
    text = '''Rating Actual* Anterior**
Fortaleza Financiera A- A-
Depósitos a Corto Plazo CP1- CP1-
Depósitos a Largo Plazo A A-
2do Programa de Certificados de Depósitos Negociables CP1- CP1-
*Información auditada al 31 de diciembre de 2025. Aprobado en comité de 27-03-2026.
**Información no auditada al 30 de junio de 2025. Aprobado en comité de 15-09-2025.
Perspectiva Estable Estable'''
    fields = extract_cover_fields(text, agency_code='001196')
    assert [(f['value_raw'],f['temporal_role']) for f in fields if f['field_kind']=='long_term_deposits'] == [('A','current'),('A-','previous')]
    assert len([f for f in fields if f['field_kind']=='short_term_deposits']) == 2
    dates = [f for f in fields if f['field_kind']=='committee_date']
    assert len(dates)==1 and dates[0]['normalized_value']=='2026-03-27'
    assert dates[0]['extraction_status']=='extracted' and '**Información' not in dates[0]['evidence_text']
    assert not any(f['field_kind']=='long_term_deposits' for f in extract_cover_fields(text,agency_code='000408'))


def test_jcr_committee_date_is_not_borrowed_from_previous_note():
    text = '''Rating Actual* Anterior**
Depósitos a Largo Plazo A A-
*Información auditada al 31 de diciembre de 2025.
**Información a junio de 2025. Aprobado en comité de 15-09-2025.
Perspectiva Estable Estable'''
    assert not any(f['field_kind']=='committee_date' for f in extract_cover_fields(text,agency_code='001196'))
