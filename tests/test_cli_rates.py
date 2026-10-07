import pandas as pd
import pytest
from openpyxl import load_workbook

from fuentes_financieras.cli import passive_rates as cli_rates
from fuentes_financieras.models import SyncResult


def test_invalid_date_range_exits_before_fetch():
    with pytest.raises(SystemExit):
        cli_rates.parse_args(['--desde', '2026-10-07', '--hasta', '2026-10-01'])


def test_excel_export_contains_table_and_filters(tmp_path):
    path = tmp_path / 'rates.xlsx'
    cli_rates.export_excel(pd.DataFrame({'Entidad': ['Banco'], 'Tasa': [4.5]}), path)
    book = load_workbook(path)
    sheet = book.active
    assert sheet['B2'].value == 4.5
    assert sheet.freeze_panes == 'A2'
    table = sheet.tables['TasasPasivas']
    assert table.ref == 'A1:B2'
    assert table.autoFilter.ref == 'A1:B2'
    assert table.tableStyleInfo.name == 'TableStyleLight9'


def test_failed_sync_does_not_export_partial_data(tmp_path, monkeypatch):
    class FailedSource:
        def sync(self, **query):
            return SyncResult('pe.sbs.tasas_pasivas', requested=2, downloaded=1, failed=1)

        def load(self, **query):
            raise AssertionError('A partial sync must not produce an apparently complete export')

    monkeypatch.setattr(cli_rates, 'source', lambda _: FailedSource())
    path = tmp_path / 'partial.xlsx'
    assert cli_rates.main(['--desde', '2026-09-01', '--excel', str(path)]) == 1
    assert not path.exists()


def test_load_only_never_calls_sync(monkeypatch):
    class CachedSource:
        def sync(self, **query):
            raise AssertionError('Network synchronization is disabled')

        def load(self, **query):
            assert query['tipos'] == ['B', 'C']
            return pd.DataFrame({'rate': [4.5]})

    monkeypatch.setattr(cli_rates, 'source', lambda _: CachedSource())
    assert cli_rates.main(['--desde', '2026-09-01', '--tipos', 'B', 'C', '--load-only']) == 0
