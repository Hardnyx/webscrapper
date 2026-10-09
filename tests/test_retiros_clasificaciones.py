"""Withdrawals need explicit scope; missing ratings and historical text do not qualify."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from pypdf import PdfWriter

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import InvalidQueryError, SchemaChangedError, SourceUnavailableError
from fuentes_financieras.providers.clasificadoras.retiros_clasificaciones import (
    RatingWithdrawalsProvider, document_reference, extract_withdrawal, parse_document,
)

URL = 'https://moodyslocal.com.pe/wp-content/uploads/2026/07/ejemplo.pdf'
AAI_URL = 'https://www.aai.com.pe/wp-content/uploads/2025/09/ejemplo.pdf'
AUTHORITY = 'Moody’s Local PE Clasificadora de Riesgo S.A. (en adelante, Moody’s Local Perú) '
LEAD = ('afirma y retira la clasificación C otorgada como Entidad a Financiera Ejemplo S.A. '
        '(en adelante, la Financiera) y la clasificación BB.pe como Emisor.')
PROGRAM = ('afirma la categoría A como Entidad a Banco Ejemplo S.A. (en adelante, el Banco). '
           'Asimismo, afirma la categoría AA.pe como Emisor y los Depósitos de Corto Plazo. '
           'De otro lado, retira la clasificación del Quinto Programa de Certificados de Depósito '
           'Negociables, debido a su vencimiento. La Perspectiva es Estable.')
AAI_LEAD = ('se pone en conocimiento público que se retiran las clasificaciónes de Caja Ejemplo '
            'de institución de B+ y de sus instrumentos: depósitos de largo plazo de A+(pe), '
            'depósitos de corto plazo de CP1 -(pe), certificados de depósitos negociables de CP1-(pe), '
            'y bonos subordinados de A(pe) por resolución de contrato.')


def cover(lead=LEAD, date='30 de julio de 2026'):
    return ('Perú COMUNICADO DE PRENSA Moody’s Local Perú afirma y retira las clasificaciones '
            'ACCIÓN DE CLASIFICACIÓN LIMA, PERÚ ' + date + ' ' + AUTHORITY + lead
            + ' La acción de clasificación se resume en el siguiente detalle: RET historial CONTACTOS')


def aai_cover(lead=AAI_LEAD):
    return ('Apoyo & Asociados retira las clasificaciones\n25.09.2025\n'
            'Apoyo & Asociados Internacionales (A&A) – Lima, 25 de setiembre de 2025: '
            + lead + ' Contactos: Otro texto histórico')


def test_entity_and_issuer_scope_is_explicit_not_all_ratings():
    row = extract_withdrawal(cover(), agency='Moody’s Local Perú')
    assert row['event_type'] == 'rating_withdrawal_announced' and row['announcement_date'] == '2026-07-30'
    assert row['entity_name'] == 'Financiera Ejemplo S.A.'
    assert row['withdrawal_scope'] == 'Entidad: C; Emisor: BB.pe' and row['withdrawal_reason'] == ''
    assert LEAD in row['evidence_text'] and 'historial' not in row['evidence_text']


def test_program_retirement_does_not_withdraw_affirmed_deposits():
    row = extract_withdrawal(cover(PROGRAM), agency='Moody’s Local Perú')
    assert row['entity_name'] == 'Banco Ejemplo S.A.'
    assert row['withdrawal_scope'] == 'Quinto Programa de Certificados de Depósito Negociables'
    assert row['withdrawal_reason'] == 'su vencimiento'
    assert 'Depósitos de Corto Plazo' not in row['withdrawal_scope']


def test_aai_preserves_alias_exact_ratings_and_contract_reason():
    row = extract_withdrawal(aai_cover(), agency='Apoyo & Asociados')
    assert row['announcement_date'] == '2025-09-25' and row['entity_name'] == 'Caja Ejemplo'
    assert row['withdrawal_reason'] == 'resolución de contrato'
    assert 'CP1 -(pe)' in row['withdrawal_scope'] and 'bonos subordinados' in row['withdrawal_scope']


@pytest.mark.parametrize('lead', [LEAD.replace('afirma y retira','afirma'),
    LEAD.replace('afirma y retira','no retira'), LEAD.replace('afirma y retira','podría retirar'),
    'Según una clasificación anterior: '+LEAD, LEAD+' En un periodo anterior se retiró otra clasificación.',
    PROGRAM.replace('retira la clasificación','no retira la clasificación'),
    PROGRAM.replace('retira la clasificación','podría retirar la clasificación')])
def test_affirmations_hypotheses_and_mixed_historical_text_need_review(lead):
    row = extract_withdrawal(cover(lead), agency='Moody’s Local Perú')
    assert row['extraction_status'] == 'needs_review' and row['event_type'] == '' and row['withdrawal_scope'] == ''


def test_title_and_table_markers_do_not_establish_withdrawal():
    row = extract_withdrawal(cover('afirma las clasificaciones vigentes.'), agency='Moody’s Local Perú')
    assert row['event_type'] == ''
    text = cover('afirma las clasificaciones vigentes.') + ' Fundamentos de la clasificación ' + LEAD
    assert extract_withdrawal(text, agency='Moody’s Local Perú')['event_type'] == ''
    assert extract_withdrawal('Moody’s Local retira RET RETIRADA', agency='Moody’s Local Perú')['event_type'] == ''


def test_wrong_agency_or_ambiguous_dated_blocks_need_review():
    assert extract_withdrawal(cover(), agency='Apoyo & Asociados')['event_type'] == ''
    assert extract_withdrawal(cover()+cover(), agency='Moody’s Local Perú')['event_type'] == ''
    assert extract_withdrawal(aai_cover()+aai_cover(), agency='Apoyo & Asociados')['event_type'] == ''


def test_aai_negation_and_unrecognized_scope_do_not_create_events():
    for lead in [AAI_LEAD.replace('se retiran','no se retiran'), AAI_LEAD.replace('bonos subordinados','otros instrumentos')]:
        assert extract_withdrawal(aai_cover(lead), agency='Apoyo & Asociados')['event_type'] == ''


def test_invalid_explicit_date_fails():
    with pytest.raises(SchemaChangedError):
        extract_withdrawal(cover(date='31 de febrero de 2026'), agency='Moody’s Local Perú')


@pytest.mark.parametrize('url', [URL.replace('https','http'), URL.replace('moodyslocal.com.pe','other.test'),
    URL+'?download=1', URL+'#fragment', URL.replace('/07/','/13/'), URL.replace('.pdf','.html'),
    URL.replace('ejemplo.pdf','folder/ejemplo.pdf'), URL.replace('ejemplo.pdf','folder%2Fejemplo.pdf')])
def test_only_official_unambiguous_pdf_urls_are_allowed(url):
    with pytest.raises(InvalidQueryError):document_reference(url)


def test_url_identity_deduplicates_equivalent_encoding():
    assert document_reference(URL)['document_id'] == document_reference(URL.replace('ejemplo','%65jemplo'))['document_id']
    assert document_reference(AAI_URL)['rating_agency'] == 'Apoyo & Asociados'


def blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def test_valid_blank_pdf_is_needs_ocr_not_withdrawal():
    data, metadata = parse_document(blank_pdf(), reference=document_reference(URL), retrieved_at='checked')
    assert data.extraction_status.iloc[0] == 'needs_ocr' and data.event_type.iloc[0] == ''
    assert metadata['page_count'] == 1 and metadata['text_pages'] == 0


@pytest.mark.parametrize('content', [b'<html>error</html>', b'%PDF-1.7 broken', b'%PDF-1.7 broken %%EOF'])
def test_non_pdf_truncation_and_corruption_fail(content):
    with pytest.raises((SourceUnavailableError, SchemaChangedError)):
        parse_document(content, reference=document_reference(URL), retrieved_at='checked')


def text_reader(monkeypatch, text):
    # Extraction grammar and storage integrity are tested independently of PDF typography.
    import pypdf
    monkeypatch.setattr(pypdf, 'PdfReader', lambda *a,**kw:SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda:text)]))


def test_pdf_cache_force_reparse_redownload_redirect_and_failed_download(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    text_reader(monkeypatch, cover())
    provider = RatingWithdrawalsProvider(CATALOG['pe.clasificadoras.retiros'])
    payload = blank_pdf()
    calls = []
    def request(*a,**kw):
        calls.append(a)
        return SimpleNamespace(content=payload,url=URL)
    monkeypatch.setattr(provider.transport,'request',request)
    assert provider.sync(urls=[URL,URL]).downloaded == 1 and len(calls) == 1
    assert provider.sync(urls=[URL]).skipped_existing == 1
    provider.sync(urls=[URL],force=True)
    assert len(calls) == 1
    provider.sync(urls=[URL],redownload=True)
    assert len(calls) == 2
    previous = provider.load().copy()
    path = provider.pdf_path(document_reference(URL)['document_id'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(content=payload,url=URL.replace('ejemplo','otro')))
    assert provider.sync(urls=[URL],redownload=True).failed == 1 and provider.load().equals(previous)
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(content=b'<html>error</html>',url=URL))
    assert provider.sync(urls=[URL],redownload=True).failed == 1 and provider.load().equals(previous)
    assert path.read_bytes() == payload
    path.write_bytes(b'damaged')
    monkeypatch.setattr(provider.transport,'request',request)
    restored = provider.sync(urls=[URL])
    assert restored.failed == 0 and restored.unavailable == 0 and len(calls) == 3
    assert path.read_bytes() == payload


def test_cli_offline_style_and_pdf_integrity_rejection(tmp_path, monkeypatch):
    from fuentes_financieras.cli import retiros_clasificaciones as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    text_reader(monkeypatch, cover())
    provider = RatingWithdrawalsProvider(CATALOG['pe.clasificadoras.retiros'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(content=blank_pdf(),url=URL))
    assert provider.sync(urls=[URL]).downloaded == 1
    monkeypatch.setattr(cli,'source',lambda _:provider)
    def no_network(*a,**kw):raise AssertionError('Unexpected offline request')
    monkeypatch.setattr(provider.transport,'request',no_network)
    args = ['--urls',URL,'--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args) == 0
    path = tmp_path/'reports/retiros_clasificaciones.xlsx'
    book = load_workbook(path)
    for sheet in book:
        assert sheet.freeze_panes == 'A2' and not sheet.column_dimensions
        assert next(iter(sheet.tables.values())).tableStyleInfo.name == 'TableStyleLight9'
    path.unlink()
    provider.pdf_path(document_reference(URL)['document_id']).write_bytes(b'damaged')
    assert cli.main(args) == 1 and not path.exists()


def test_cli_review_selection_has_empty_withdrawal_sheet(tmp_path, monkeypatch):
    from fuentes_financieras.cli import retiros_clasificaciones as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    text_reader(monkeypatch, cover('afirma las clasificaciones vigentes.'))
    provider = RatingWithdrawalsProvider(CATALOG['pe.clasificadoras.retiros'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(content=blank_pdf(),url=URL))
    assert provider.sync(urls=[URL]).downloaded == 1
    monkeypatch.setattr(cli,'source',lambda _:provider)
    assert cli.main(['--urls',URL,'--load-only','--output-dir',str(tmp_path/'reports')]) == 2
    book = load_workbook(tmp_path/'reports/retiros_clasificaciones.xlsx')
    assert book['retiros'].max_row == 1 and book['comunicados'].max_row == 2
