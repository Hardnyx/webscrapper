"""Guard currencies, system concentration scope, dates and monthly flow bases."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError, InvalidQueryError
from fuentes_financieras.providers.sbs.fondeo import (
    PERSON_LABELS, DEPOSIT_LABELS, TERM_LABELS, SCALE_LABELS, parse_person,
    parse_terms, parse_scale, parse_debt, DepositsByPersonProvider, DepositsByTermProvider,
)
from fuentes_financieras.providers.sbs.castigos import parse_workbook


def save(book):
    stream=BytesIO(); book.save(stream); return stream.getvalue()


def book(title, unit='(En miles de soles)', date=datetime(2026,8,31)):
    b=Workbook(); s=b.active; s.cell(1,1,title);s.cell(2,1,date);s.cell(3,1,unit)
    return b,s


def persons(*, placeholder=False, dash=False, mismatch=False, unknown=False):
    b,s=book('Depósitos por Tipo, Persona y Empresa Bancaria');s.cell(5,1,'Empresas')
    for i,label in enumerate(DEPOSIT_LABELS):
        c=2+i*3;s.cell(5,c,label)
        for k,p in enumerate(PERSON_LABELS):s.cell(6,c+k,p)
    if unknown:s.cell(6,3,'Persona no revisada')
    for r,name in ((8,0 if placeholder else 'Entidad'),(9,'TOTAL SISTEMA')):
        s.cell(r,1,name)
        for c in range(2,17):s.cell(r,c,10 if c<14 else 40)
        if dash:s.cell(r,2,'-')
        if mismatch:s.cell(r,14,50)
    s.cell(11,1,'Fuente: Anexo 13');return save(b)


def terms(*, wrong_unit=False, date=datetime(2026,8,1), missing_currency=False):
    b=Workbook();b.remove(b.active)
    for cur,unit in [('Nacional','soles'),('Extranjera','dólares')]:
        if missing_currency and cur=='Extranjera':continue
        s=b.create_sheet(cur);s.cell(1,1,'Depósitos del público en Moneda '+cur+' por Empresa Bancaria')
        s.cell(2,1,date);s.cell(3,1,'(En miles de '+('soles' if wrong_unit else unit)+')');s.cell(5,4,'Cuentas a Plazo')
        for c,label in enumerate(TERM_LABELS,2):s.cell(6,c,label)
        for r,name in [(8,'Entidad'),(9,'TOTAL SISTEMA')]:
            s.cell(r,1,name)
            for c in range(2,11):s.cell(r,c,10 if c<10 else 80)
        s.cell(11,1,'Fuente: Reporte N° 6-B')
    return save(b)


def scales(*, fractional=False, discontinuous=False, unknown=False):
    b,s=book('Depósitos de la Banca Múltiple según Escala de Montos',unit='')
    s.cell(2,1,None);s['A2']=None;s.cell(2,2,datetime(2026,8,31))
    s.cell(5,1,'Escala');s.cell(5,6,'Personas Naturales');s.cell(5,10,'Personas Jurídicas');s.cell(5,18,'TOTAL')
    s.cell(6,10,'Privadas sin fines de lucro');s.cell(6,14,'Otras personas jurídicas');s.cell(7,1,'( En soles )')
    for c in (6,10,14,18):
        s.cell(7,c,'Número');s.cell(7,c+2,'Monto');s.cell(8,c+2,'(Miles de soles)')
    for p,label in enumerate(SCALE_LABELS):
        r=10+p*4;s.cell(r,1,label)
        s.cell(r+1,2,'Hasta');s.cell(r+1,4,100)
        s.cell(r+2,1,'de');s.cell(r+2,2,200 if discontinuous else 100);s.cell(r+2,3,'a');s.cell(r+2,4,'más')
        for c in (6,10,14,18):
            s.cell(r,c,5);s.cell(r,c+2,10)
            s.cell(r+1,c,2.5 if fractional else 2);s.cell(r+1,c+2,3)
            s.cell(r+2,c,3);s.cell(r+2,c+2,7)
    if unknown:s.cell(5,18,'Entidad X')
    s.cell(31,1,'Fuente: Anexo N°13');return save(b)


def debt(*, percent_format=False):
    b,s=book('Estructura de los Adeudos y Obligaciones Financieras por Empresa Bancaria','(En porcentaje)')
    s.cell(5,1,'Empresas');s.cell(5,2,'Instituciones del País');s.cell(5,4,'Instituciones del Exterior y Organismos Internacionales')
    s.cell(5,6,'Total Adeudos y Obligaciones Financieras (En miles de soles)')
    for c,label in enumerate(['Corto Plazo','Largo Plazo']*2,2):s.cell(6,c,label)
    for r,name in [(8,'Entidad'),(9,'TOTAL SISTEMA')]:
        s.cell(r,1,name)
        for c,value in enumerate([25,25,25,25,1000],2):
            s.cell(r,c,value)
            if percent_format and c<6:s.cell(r,c).number_format='0%'
    s.cell(11,1,'Fuente: Balance de Comprobación.');return save(b)


def writeoffs(*, accumulated=False, wrong_month=False, dash=True):
    b,s=book('Flujo de Créditos Castigados por Tipo de Crédito y Empresa Bancaria',date='en el mes de Julio de 2026' if wrong_month else 'en el mes de Agosto de 2026')
    s.cell(5,1,'Empresas');s.cell(5,3,'Flujo acumulado de castigos' if accumulated else 'Flujo de castigos');s.cell(5,10,'Total')
    for c,label in enumerate(['Corporativos','Grandes Empresas','Medianas Empresas','Pequeñas Empresas','Microempresas','Consumo','Hipotecarios'],3):s.cell(6,c,label)
    for r,name in [(8,'Entidad'),(9,'TOTAL SISTEMA')]:
        s.cell(r,1,name)
        for c in range(3,11):s.cell(r,c,'-' if dash and c==3 else 10 if c<10 else 70)
    s.cell(11,1,'Fuente: Reporte N°25');s.cell(12,1,'Nota: Criterios empresariales no comparables antes de octubre de 2024.')
    return save(b)


def parse(content, parser):
    return parser(content,entity_type='B',period='2026-08',source_url='test',retrieved_at='test')


def test_persons_preserve_unknown_entity_and_dash_without_zero_inference():
    d,_=parse(persons(placeholder=True,dash=True),parse_person)
    assert set(d.currency)=={'TOTAL'} and set(d.unit)=={'thousands_PEN'}
    unknown=d[d.entity_scope=='unidentified_source_row']
    assert len(unknown)==15 and set(unknown.entity_name)=={''} and set(unknown.source_entity_name)=={'0'}
    assert d.value.isna().sum()==2 and set(d[d.value.isna()].source_value_token)=={'-'}
    mismatch,_=parse(persons(mismatch=True),parse_person)
    assert 'published_components_mismatch' in set(mismatch.data_quality_flags)
    with pytest.raises(SchemaChangedError):parse(persons(unknown=True),parse_person)


def test_terms_keep_published_date_and_currency_units(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    d,notes=parse(terms(),parse_terms)
    assert set(d.period_date)=={'2026-08-01'}
    assert set(d.data_quality_flags)=={'source_date_not_month_end'}
    assert set(d[d.currency=='ME'].unit)=={'thousands_USD'}
    assert set(d[d.currency=='MN'].unit)=={'thousands_PEN'}
    assert set(d[d.deposit_type=='term'].term_bucket)=={'le_30','31_90','91_180','181_360','gt_360'}
    assert len(notes)==4
    for kw in [{'wrong_unit':True},{'missing_currency':True},{'date':datetime(2026,7,31)}]:
        with pytest.raises(SchemaChangedError):parse(terms(**kw),parse_terms)
    p=DepositsByTermProvider(CATALOG['pe.sbs.depositos_plazo'])
    assert {r.params['tipo'] for r in p.plan_sync(desde='2026-08')}=={'B','F'}
    with pytest.raises(InvalidQueryError):list(p.plan_sync(desde='2026-08',tipos=['C']))


def test_scales_remain_system_aggregate_with_unknown_boundary_inclusivity():
    d,_=parse(scales(),parse_scale)
    assert len(d)==120 and set(d.entity_scope)=={'system_aggregate'}
    assert set(d.metric)=={'published_number','deposit_amount'}
    assert set(d.boundary_convention)=={'unspecified_by_source'}
    assert d[d.band_kind=='upper_bounded'].band_lower_PEN.isna().all()
    assert d[d.band_kind=='upper_open'].band_upper_PEN.isna().all()
    assert set(d[d.band_kind=='upper_open'].band_lower_PEN)=={100}
    assert set(d.data_quality_flags)=={''}
    for kw in [{'fractional':True},{'discontinuous':True},{'unknown':True}]:
        with pytest.raises(SchemaChangedError):parse(scales(**kw),parse_scale)


def test_debt_separates_percentage_points_from_monetary_total():
    d,_=parse(debt(),parse_debt)
    assert set(d[d.metric=='financial_obligations_share'].value)=={25}
    assert set(d[d.metric=='financial_obligations_share'].unit)=={'percent'}
    assert set(d[d.metric=='financial_obligations_total'].unit)=={'thousands_PEN'}
    with pytest.raises(SchemaChangedError):parse(debt(percent_format=True),parse_debt)


def test_writeoffs_monthly_window_and_classification_notes():
    d,notes=parse(writeoffs(),parse_workbook)
    assert set(d.measurement_basis)=={'monthly_flow'}
    assert set(d.observation_start)=={'2026-08-01'} and set(d.observation_end)=={'2026-08-31'}
    assert d.value.isna().sum()==2 and set(d[d.value.isna()].source_value_token)=={'-'}
    assert any('octubre de 2024' in n.get('text','') for n in notes)
    for kw in [{'accumulated':True},{'wrong_month':True}]:
        with pytest.raises(SchemaChangedError):parse(writeoffs(**kw),parse_workbook)


def test_funding_cache_skips_requests_and_missing_month_is_unavailable(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=DepositsByPersonProvider(CATALOG['pe.sbs.depositos_persona']);calls=[]
    url='https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-2372-ag2026.XLS'
    def request(method,target):
        calls.append(target);return SimpleNamespace(text=f'<a href="{url}">Agosto</a>',content=persons())
    monkeypatch.setattr(p.transport,'request',request)
    assert p.sync(desde='2026-08',tipos=['B']).downloaded==1
    assert p.sync(desde='2026-08',tipos=['B']).skipped_existing==1 and len(calls)==2
    assert p.sync(desde='2026-09',tipos=['B']).unavailable==1


def test_funding_cli_rejects_partial_cache_and_preserves_spanish_report(tmp_path,monkeypatch):
    from fuentes_financieras.cli import fondeo
    from fuentes_financieras.models import SyncResult
    d,_=parse(persons(),parse_person)
    request=SimpleNamespace(params={'tipo':'B','periodo':'2026-08'})
    fake=SimpleNamespace(plan_sync=lambda **kw:[request],sync=lambda **kw:SyncResult('test',requested=1,downloaded=1),load=lambda **kw:d)
    monkeypatch.setattr(fondeo,'source',lambda dataset:fake)
    args=['--datasets','personas','--desde','2026-08','--tipos','B','--output-dir',str(tmp_path)]
    assert fondeo.main(args)==0
    b=load_workbook(tmp_path/'fondeo.xlsx');s=b.active
    assert s.freeze_panes=='A2' and next(iter(s.tables.values())).tableStyleInfo.name=='TableStyleLight9'
    assert 'Tipo de depósito' in [c.value for c in s[1]]
    (tmp_path/'fondeo.xlsx').unlink()
    fake.sync=lambda **kw:SyncResult('test',requested=1,unavailable=1)
    assert fondeo.main(args)==1 and not (tmp_path/'fondeo.xlsx').exists()
    fake.load=lambda **kw:d.iloc[:0]
    assert fondeo.main(args+['--load-only'])==1
