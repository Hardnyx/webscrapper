"""Verify source dates, currency units and Excel percentage semantics."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.providers.sbs.liquidez import (
    LiquidityProvider, parse_liquidity, parse_coverage, parse_funding,
)


def workbook(kind='liquidity', *, date='2026-08-31', ratio_format='0%', currency_unit='dólares', ratio=1.25, coverage_format='General'):
    book = Workbook(); s = book.active
    if kind == 'liquidity':
        s.cell(1, 1, 'Ratios de Liquidez en Moneda Nacional y Moneda Extranjera por Empresa Bancaria')
        s.cell(2, 1, datetime.fromisoformat(date)); s.cell(4, 1, 'Empresas')
        for first, currency, unit in ((2, 'Nacional', 'soles'), (6, 'Extranjera', currency_unit)):
            s.cell(4, first, f'Liquidez en Moneda {currency} (En miles de {unit})')
            for col, label in ((first, 'Activos Líquidos'), (first+1, 'Pasivos de Corto Plazo'), (first+2, 'Ratio de Liquidez (En porcentaje)')):
                s.cell(5, col, label)
            s.cell(7, first, 50); s.cell(7, first+1, 200); s.cell(7, first+2, 25)
            # A published blank must stay missing rather than becoming zero.
            s.cell(8, first+1, 200); s.cell(8, first+2, 25)
        s.cell(7, 1, 'Entidad'); s.cell(8, 1, 'TOTAL SISTEMA'); s.cell(10, 1, 'NOTA: Anexo 15-C')
    elif kind == 'coverage':
        s.cell(2, 2, 'Ratio de Cobertura de Liquidez')
        s.cell(3, 2, 'Saldos y ratio promedio diario de enero a marzo de 2026 (1)')
        s.cell(6, 5, 'CONSOLIDADO SISTEMA')
        for col, group in ((5, 'MONEDA NACIONAL (En miles de soles)'), (7, 'MONEDA EXTRANJERA (En miles de dólares)'), (9, 'TOTAL (En miles de soles)')):
            s.cell(7, col, group); s.cell(8, col, 'Importe Base (promedio)'); s.cell(8, col+1, 'Importe Ajustado (promedio)')
        for row, label in ((32, 'Total ALAC'), (33, 'Total Flujos Entrantes 30 días'), (34, 'Total Flujos Salientes 30 días'), (35, 'RATIO DE COBERTURA DE LIQUIDEZ (%) (2)')):
            s.cell(row, 3, label)
            for col in range(5, 11):
                if row < 35 or col in (6, 8, 10):
                    s.cell(row, col, 180 if row == 35 else 100)
                    if row == 35:s.cell(row, col).number_format = coverage_format
        s.cell(36, 2, 'Fuente: Anexo 15-B'); s.cell(40, 2, 'Promedio de los ratios diarios del trimestre')
    else:
        s.cell(2, 2, 'Ratio de Financiación Neta Estable'); s.cell(3, 2, datetime.fromisoformat(date))
        s.cell(5, 4, 'CONSOLIDADO SISTEMA')
        s.cell(6, 4, 'VALOR NO PONDERADO POR VENCIMIENTO RESIDUAL (En miles de Soles)')
        s.cell(6, 8, 'VALOR PONDERADO (En miles de Soles)')
        for row, label, value in ((53, 'Total Financiación Estable Disponible', 125), (54, 'Total Financiación Estable Requerida', 100), (55, 'RATIO DE FINANCIACIÓN NETA ESTABLE (%)', ratio)):
            s.cell(row, 3, label); s.cell(row, 8, value)
        s.cell(55, 8).number_format = ratio_format
        s.cell(56, 2, 'Fuente: Anexo 16-C'); s.cell(61, 2, 'División de financiación disponible entre requerida')
    buffer = BytesIO(); book.save(buffer)
    return buffer.getvalue()


def parse(kind, **kwargs):
    return {'liquidity': parse_liquidity, 'coverage': parse_coverage, 'funding': parse_funding}[kind](
        workbook(kind, **kwargs), entity_type='B', period='2026-06' if kind == 'coverage' else '2026-08', source_url='test', retrieved_at='test')[0]


def test_currency_scale_and_missing_values():
    data = parse('liquidity')
    assert set(data[data.currency == 'USD'].unit) == {'thousands_USD', 'percent'}
    assert data.value.isna().sum() == 2
    assert set(data[data.value.isna()].data_quality_flags) == {'source_value_missing'}
    assert set(data[data.metric == 'liquidity_ratio'].value) == {25}
    assert set(data.period_date) == {'2026-08-31'}


@pytest.mark.parametrize('kwargs', [{'date': '2026-07-31'}, {'currency_unit': 'euros'}])
def test_liquidity_rejects_changed_dates_and_units(kwargs):
    with pytest.raises(SchemaChangedError):parse('liquidity', **kwargs)


def test_coverage_uses_declared_quarter_instead_of_link_month():
    data = parse('coverage')
    assert len(data) == 21
    assert set(data.period) == {'2026-06'}
    assert set(data.observation_start) == {'2026-01-01'}
    assert set(data.period_date) == {'2026-03-31'}
    assert set(data.frequency) == {'quarterly'}
    assert set(data.data_quality_flags) == {'source_period_differs_from_index'}
    # The average of daily ratios is published independently of average balances.
    assert set(data[data.metric == 'liquidity_coverage_ratio'].value) == {180}


def test_coverage_rejects_changed_ratio_scale_or_future_quarter():
    with pytest.raises(SchemaChangedError):parse('coverage', coverage_format='0%')
    with pytest.raises(SchemaChangedError):
        parse_coverage(workbook('coverage'), entity_type='B', period='2026-02', source_url='test', retrieved_at='test')


def test_funding_converts_only_verified_excel_percentage_format():
    data = parse('funding')
    ratio = data[data.metric == 'net_stable_funding_ratio'].iloc[0]
    assert ratio.value == 125 and ratio.source_value == 1.25
    assert ratio.unit == 'percent' and ratio.source_number_format == '0%'
    assert set(data[data.metric != 'net_stable_funding_ratio'].unit_multiplier) == {1000}
    with pytest.raises(SchemaChangedError):parse('funding', ratio_format='General')
    with pytest.raises(SchemaChangedError):parse('funding', date='2026-07-31')


def test_preserves_contradictory_funding_ratio():
    data = parse('funding', ratio=2)
    assert set(data.data_quality_flags) == {'published_ratio_mismatch'}
    assert data[data.metric == 'net_stable_funding_ratio'].value.iloc[0] == 200


def test_cache_skips_network_and_missing_month_is_unavailable(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path))
    provider = LiquidityProvider(CATALOG['pe.sbs.liquidez']); calls = []
    url = 'https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-2340-ag2026.XLS'
    def request(method, target):
        calls.append(target)
        return SimpleNamespace(text=f'<a href="{url}">Agosto</a>', content=workbook())
    monkeypatch.setattr(provider.transport, 'request', request)
    assert provider.sync(desde='2026-08', tipos=['B']).downloaded == 1
    assert provider.sync(desde='2026-08', tipos=['B']).skipped_existing == 1
    assert len(calls) == 2
    assert provider.sync(desde='2026-09', tipos=['B']).unavailable == 1


def test_cli_does_not_export_incomplete_capture(tmp_path, monkeypatch):
    from fuentes_financieras.cli import liquidez
    fake = SimpleNamespace(plan_sync=lambda **kw: [], sync=lambda **kw: SimpleNamespace(failed=0, unavailable=1))
    monkeypatch.setattr(liquidez, 'source', lambda dataset: fake)
    assert liquidez.main(['--desde', '2026-08', '--output-dir', str(tmp_path)]) == 1
    assert not (tmp_path / 'liquidez.xlsx').exists()
