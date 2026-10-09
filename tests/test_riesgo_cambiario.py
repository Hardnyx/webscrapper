"""Guard signed PEN exposures, missing components and lagged capital dates."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
import pytest
from openpyxl import Workbook, load_workbook
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.models import SyncResult
from fuentes_financieras.providers.sbs.riesgo_cambiario import (
    GROUP_NAMES, POSITION_LABELS, CAPITAL_NOTE, RATIO_LABEL,
    parse_positions, parse_capital_ratio, ForeignExchangePositionProvider,
)


def book(*, ratio=False, tipo='B', date=datetime(2026,7,31), missing=False,
         mismatch=False, unit_changed=False, note=CAPITAL_NOTE, formatted=False):
    b=Workbook();s=b.active
    if ratio:
        for col,name in [(1,'Entidad'),(4,'TOTAL SISTEMA')]:
            s.cell(1,col,'Indicadores Financieros por '+GROUP_NAMES[tipo])
            s.cell(2,col,date);s.cell(4,col+1,name)
            s.cell(7,col,'POSICIÓN EN MONEDA EXTRANJERA')
            s.cell(8,col,RATIO_LABEL+('' if unit_changed else ' ( %) ***'))
            s.cell(8,col+1,-125)
            if formatted:s.cell(8,col+1).number_format='0%'
            s.cell(9,col,'Nota: Definiciones originales')
            if note:s.cell(10,col,'*** '+note)
    else:
        s.cell(1,1,'Posición Global en Moneda Extranjera por '+GROUP_NAMES[tipo])
        s.cell(2,1,date);s.cell(3,1,'(En miles de dólares)' if unit_changed else '(En miles de soles)')
        s.cell(5,1,'Empresas')
        for col,label in enumerate(POSITION_LABELS,2):s.cell(5,col,label)
        for row,name in [(7,'Entidad'),(8,'TOTAL SISTEMA')]:
            s.cell(row,1,name)
            for col,value in enumerate([-10, '-' if missing else 5, -1, 9 if mismatch else -6],2):s.cell(row,col,value)
        s.cell(10,1,'Fuente: Reporte original SBS')
    stream=BytesIO();b.save(stream);return stream.getvalue()


def parse(content, *, ratio=False, tipo='B', period='2026-07'):
    return (parse_capital_ratio if ratio else parse_positions)(content,entity_type=tipo,
        period=period,source_url='test',retrieved_at='test')[0]


@pytest.mark.parametrize('tipo',list('BFCR'))
def test_position_keeps_signed_pen_values_and_does_not_zero_missing_components(tipo):
    d=parse(book(tipo=tipo,missing=True),tipo=tipo)
    assert len(d)==8 and set(d.unit)=={'thousands_PEN'} and set(d.unit_multiplier)=={1000}
    assert set(d.currency)=={'ME'} and set(d.denominator_period)=={''}
    assert d.value.isna().sum()==2 and set(d[d.value.isna()].source_value_token)=={'-'}
    assert set(d[d.metric=='global_fx_position'].value)=={-6}
    assert not d.data_quality_flags.str.contains('published_components_mismatch').any()


def test_published_component_mismatch_is_flagged_without_replacing_values():
    d=parse(book(mismatch=True))
    assert set(d.data_quality_flags)=={'published_components_mismatch'}
    assert set(d[d.metric=='global_fx_position'].value)=={9}


@pytest.mark.parametrize('kwargs',[{'unit_changed':True},{'date':datetime(2026,7,30)}])
def test_position_rejects_changed_units_or_date(kwargs):
    with pytest.raises(SchemaChangedError):parse(book(**kwargs))


@pytest.mark.parametrize('tipo',['C','R'])
def test_capital_ratio_preserves_sign_and_previous_year_denominator(tipo):
    d=parse(book(ratio=True,tipo=tipo,date=datetime(2026,1,31)),ratio=True,tipo=tipo,period='2026-01')
    assert len(d)==2 and set(d.value)=={-125} and set(d.unit)=={'percent'}
    assert set(d.denominator_period)=={'2025-12'} and set(d.denominator_date)=={'2025-12-31'}
    assert set(d.period_date)=={'2026-01-31'} and set(d.data_quality_flags)=={''}


@pytest.mark.parametrize('kwargs',[{'note':''},{'note':'Patrimonio del mes actual.'},{'unit_changed':True},{'formatted':True}])
def test_ratio_rejects_missing_lag_evidence_and_percentage_scale_changes(kwargs):
    with pytest.raises(SchemaChangedError):parse(book(ratio=True,tipo='C',**kwargs),ratio=True,tipo='C')


def test_entity_type_cannot_be_inferred_from_wrong_workbook():
    with pytest.raises(SchemaChangedError):parse(book(tipo='C'),tipo='B')
    with pytest.raises(SchemaChangedError):parse(book(ratio=True,tipo='C'),ratio=True,tipo='R')


def test_position_cache_skips_network_and_preserves_valid_capture_on_failure(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    p=ForeignExchangePositionProvider(CATALOG['pe.sbs.posicion_cambiaria']);calls=[]
    url='https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Julio/B-2368-jl2026.XLS'
    def request(method,target):
        calls.append(target);return SimpleNamespace(text=f'<a href="{url}">Julio</a>',content=book())
    monkeypatch.setattr(p.transport,'request',request)
    assert p.sync(desde='2026-07',tipos=['B']).downloaded==1
    assert p.sync(desde='2026-07',tipos=['B']).skipped_existing==1 and len(calls)==2
    previous=p.load().copy()
    monkeypatch.setattr(p.transport,'request',lambda *a,**kw:SimpleNamespace(text=f'<a href="{url}">Julio</a>',content=book(unit_changed=True)))
    assert p.sync(desde='2026-07',tipos=['B'],force=True).failed==1 and p.load().equals(previous)
    assert p.sync(desde='2026-08',tipos=['B']).unavailable==1


def test_cli_default_position_spanish_report_and_explicit_capital_coverage(tmp_path,monkeypatch):
    from fuentes_financieras.cli import fondeo, riesgo_cambiario
    d=parse(book());selected=[]
    fake=SimpleNamespace(plan_sync=lambda **kw:[SimpleNamespace(params={'tipo':'B','periodo':'2026-07'})],
        sync=lambda **kw:SyncResult('test',requested=1,downloaded=1),load=lambda **kw:d)
    def source(dataset):selected.append(dataset);return fake
    monkeypatch.setattr(fondeo,'source',source)
    args=['--desde','2026-07','--tipos','B','--output-dir',str(tmp_path)]
    assert riesgo_cambiario.main(args)==0 and selected==['pe.sbs.posicion_cambiaria']
    path=tmp_path/'riesgo_cambiario.xlsx';s=load_workbook(path).active
    assert s.freeze_panes=='A2' and not s.column_dimensions
    assert next(iter(s.tables.values())).tableStyleInfo.name=='TableStyleLight9'
    assert 'Mes del denominador' in [c.value for c in s[1]]
    path.unlink();fake.sync=lambda **kw:SyncResult('test',requested=1,unavailable=1)
    assert riesgo_cambiario.main(args)==1 and not path.exists()
    with pytest.raises(SystemExit):riesgo_cambiario.main(args+['--datasets','capital'])
