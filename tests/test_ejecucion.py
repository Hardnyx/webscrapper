"""Verify bounded batch plans, independent failures and honest cache coverage."""
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from fuentes_financieras import ejecucion
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import PeriodUnavailableError
from fuentes_financieras.models import FetchResult, PeriodRequest
from fuentes_financieras.provider import DatasetProvider


class SampleProvider(DatasetProvider):
    def __init__(self, spec):
        super().__init__(spec)
        self.calls = 0
        self.fail = False
        self.unavailable = False

    def single_request(self, **query):
        return PeriodRequest(query.get('period', '2026-01'), 'year=2026', {})

    def plan_sync(self, **query):
        for period in query.get('periods', ['2026-01']):
            yield self.single_request(period=period)

    def _fetch_period(self, request):
        self.calls += 1
        if self.fail:
            raise RuntimeError('Controlled capture failure')
        if self.unavailable:
            raise PeriodUnavailableError('Not published')
        return FetchResult(self.spec.dataset_id, pd.DataFrame({'value': pd.Series([1], dtype='int64')}))


def setup_providers(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path/'cache'))
    providers = {key: SampleProvider(CATALOG[key]) for key in ('pe.sbs.rentabilidad', 'pe.sbs.participacion')}
    monkeypatch.setattr(ejecucion, 'get_provider', providers.__getitem__)
    return providers


def job(name='first', dataset='pe.sbs.rentabilidad', **query):
    return {'nombre': name, 'dataset': dataset, 'consulta': query}


def test_batch_reuses_real_canonical_cache_and_repeated_runs_skip(tmp_path, monkeypatch):
    providers = setup_providers(tmp_path, monkeypatch)
    plan = ejecucion.plan_jobs([job(periods=['2026-01','2026-02'])])
    summary, coverage, code = ejecucion.execute_jobs(plan)
    assert code == 0 and summary[0]['Descargadas'] == 2
    assert all(c['Estado de captura'] == 'Captura validada' for c in coverage)
    summary, _, code = ejecucion.execute_jobs(plan)
    assert code == 0 and summary[0]['Omitidas por caché'] == 2
    assert providers['pe.sbs.rentabilidad'].calls == 2


def test_plan_limit_or_invalid_later_job_cancels_all_downloads(tmp_path, monkeypatch):
    providers = setup_providers(tmp_path, monkeypatch)
    with pytest.raises(ValueError):
        ejecucion.plan_jobs([job(periods=['a','b']),job('second')], max_requests=2)
    with pytest.raises(ValueError):
        ejecucion.plan_jobs([job(),job('second',periods=[])])
    with pytest.raises(ValueError):
        ejecucion.plan_jobs([job(periods=['a','a'])])
    assert all(p.calls == 0 for p in providers.values())


def test_failed_refresh_keeps_old_capture_but_does_not_claim_success(tmp_path, monkeypatch):
    providers = setup_providers(tmp_path, monkeypatch)
    first = providers['pe.sbs.rentabilidad']
    ejecucion.execute_jobs(ejecucion.plan_jobs([job()]))
    first.fail = True
    one = job(); one['opciones'] = {'force': True}
    summary, coverage, code = ejecucion.execute_jobs(ejecucion.plan_jobs([one,job('second','pe.sbs.participacion')]))
    assert code == 1 and summary[0]['Estado'] == 'Falló'
    assert summary[1]['Estado'] == 'Completo para la selección'
    assert coverage[0]['Estado de captura'].startswith('Falló actualización')
    assert first.storage.read_all().value.tolist() == [1]


def test_unavailable_cached_period_remains_pending(tmp_path, monkeypatch):
    p = setup_providers(tmp_path, monkeypatch)['pe.sbs.rentabilidad']; p.unavailable = True
    planned = ejecucion.plan_jobs([job()])
    assert ejecucion.execute_jobs(planned)[2] == 2
    summary, coverage, code = ejecucion.execute_jobs(planned)
    assert code == 2 and summary[0]['No disponibles'] == 0
    assert coverage[0]['Estado de captura'] == 'No disponible según la fuente'
    assert p.calls == 1


def test_cache_only_detects_parser_change_corruption_and_missing_without_network(tmp_path, monkeypatch):
    p = setup_providers(tmp_path, monkeypatch)['pe.sbs.rentabilidad']
    planned = ejecucion.plan_jobs([job()]); ejecucion.execute_jobs(planned)
    p.parser_version = 'new'
    assert ejecucion.execute_jobs(planned, cache_only=True)[1][0]['Estado de captura'].startswith('Parser desactualizado')
    p.parser_version = '1'
    p.storage.partition_path('year=2026').write_bytes(b'broken parquet')
    assert ejecucion.execute_jobs(planned, cache_only=True)[2] == 2
    absent = ejecucion.plan_jobs([job(periods=['2026-03'])])
    assert ejecucion.execute_jobs(absent, cache_only=True)[1][0]['Estado de captura'] == 'Ausente'
    assert p.calls == 1


def test_plan_only_never_fetches(tmp_path, monkeypatch):
    p = setup_providers(tmp_path, monkeypatch)['pe.sbs.rentabilidad']
    summary, coverage, code = ejecucion.execute_jobs(ejecucion.plan_jobs([job()]), plan_only=True)
    assert code == 0 and summary[0]['Estado'] == 'Planeado' and coverage[0]['Estado de captura'] == 'Planeada'
    assert p.calls == 0


@pytest.mark.parametrize('body', [
    'version=2\ntrabajos=[]', 'version=1\ntrabajos=[]',
    'version=1\n[[trabajos]]\nnombre="x"\ndataset="unknown"',
    'version=1\n[[trabajos]]\nnombre="x"\ndataset="pe.sbs.rentabilidad"\n[trabajos.opciones]\nallow_schema_change=true',
    'version=1\n[[trabajos]]\nnombre="x"\ndataset="pe.sbs.rentabilidad"\n[trabajos.opciones]\nrefresh_hours=nan',
    'version=1\n[[trabajos]]\nnombre="x"\ndataset="pe.sbs.rentabilidad"\n[trabajos.opciones]\nforce="true"',
    'version=1\n[[trabajos]]\nnombre="x"\ndataset="pe.sbs.rentabilidad"\n[trabajos.consulta]\nforce=true',
])
def test_invalid_profiles_rejected_before_execution(tmp_path, body):
    path = tmp_path/'profile.toml';path.write_text(body)
    with pytest.raises(Exception):
        ejecucion.read_profile(path)


def test_cli_exports_machine_status_and_spanish_excel_without_overwriting(tmp_path,monkeypatch):
    from fuentes_financieras.cli import ejecucion as cli
    providers = setup_providers(tmp_path,monkeypatch)
    profile = tmp_path/'profile.toml'
    profile.write_text('version=1\n[[trabajos]]\nnombre="Rentabilidad"\ndataset="pe.sbs.rentabilidad"')
    args = ['--perfil',str(profile),'--data-root',str(tmp_path/'cache'),'--output-dir',str(tmp_path/'reports')]
    assert cli.run(args) == 0
    assert cli.run(args+['--solo-cache']) == 0
    runs = list((tmp_path/'reports').iterdir());assert len(runs) == 2
    for run in runs:
        payload = json.loads((run/'ejecucion.json').read_text())
        assert payload['codigo_salida'] == 0
        wb = load_workbook(run/'ejecucion.xlsx')
        assert wb.sheetnames == ['trabajos','capturas']
        for ws in wb:
            assert ws.freeze_panes == 'A2' and not ws.column_dimensions
            assert all(t.tableStyleInfo.name == 'TableStyleLight9' for t in ws.tables.values())
    assert providers['pe.sbs.rentabilidad'].calls == 1


def test_ratings_offline_plan_does_not_discover_publications(tmp_path,monkeypatch):
    from fuentes_financieras.providers.sbs.clasificaciones_riesgo import RiskRatingsProvider
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path/'cache'))
    provider = RiskRatingsProvider(CATALOG['pe.sbs.clasificaciones_riesgo'])
    def deny():
        raise AssertionError('Online discovery forbidden')
    monkeypatch.setattr(provider,'available_periods',deny)
    requests = list(provider.plan_offline(periodos=['202601','202502']))
    assert [r.period_key for r in requests] == ['202601','202502']
    assert all(r.mutable for r in requests)
    with pytest.raises(Exception):
        list(provider.plan_offline(desde='2025'))


def test_pdf_integrity_and_pending_fields_are_separate_from_capture_completeness(tmp_path,monkeypatch):
    p = setup_providers(tmp_path,monkeypatch)['pe.sbs.rentabilidad']
    planned = ejecucion.plan_jobs([job()]);ejecucion.execute_jobs(planned)
    entry = p.storage.manifest.get('2026-01')
    entry['metadata']={'field_counts':{'extracted':3,'needs_review':2}}
    p.storage.manifest.set('2026-01',entry)
    summary,coverage,code = ejecucion.execute_jobs(planned,cache_only=True)
    assert code == 0 and coverage[0]['Campos pendientes de revisión']==2
    p.pdf_matches = lambda *args: False
    assert ejecucion.execute_jobs(planned,cache_only=True)[1][0]['Estado de captura']=='PDF ausente o dañado'


def test_post_capture_freshness_does_not_reapply_redownload_trigger(tmp_path,monkeypatch):
    p = setup_providers(tmp_path,monkeypatch)['pe.sbs.rentabilidad']
    planned = ejecucion.plan_jobs([job()]);ejecucion.execute_jobs(planned)
    # A provider's refresh hook may force redownload for the current command.
    monkeypatch.setattr(p,'_refresh_due',lambda *args: True)
    assert ejecucion.execute_jobs(planned,cache_only=True)[2] == 0
