"""Keep authorizations, clarifications and unknown execution separate."""
from types import SimpleNamespace

import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import InvalidQueryError, SchemaChangedError
from fuentes_financieras.providers.elperuano.fusiones import (
    MergerResolutionsProvider, parse_resolution, resolution_reference,
)

URL = 'https://busquedas.elperuano.pe/dispositivo/NL/1234567-1'
API = 'https://busquedas.elperuano.pe/api/visor_html/1234567-1'
AUTH = ('Artículo Primero.- Autorizar la fusión por absorción de Banco Ejemplo S.A. '
        'con Financiera Ejemplo S.A., extinguiéndose esta última sin disolverse ni liquidarse.')
TITLE = 'Autorizan la fusión por absorción de Banco Ejemplo S.A. con Financiera Ejemplo S.A.'
CLARIFY = ('Artículo Único.- Autorizar a Banco Ejemplo S.A., la aprobación de la aclaratoria '
           'de la minuta de fusión por absorción con Financiera Ejemplo S.A. relativa a la fecha '
           'de entrada en vigencia de dicha fusión, estipulada el 01 de enero de 2022, así como, '
           'autorizar la modificación del Estatuto social de la entidad.')


def html(article=AUTH, title=TITLE, before='', after=''):
    return (f'<html><head><title>Título ajeno del cuadernillo</title></head><body>'
            f'<div id="x1234567-1"><div class="story"><h1 class="sumilla">{title}</h1>'
            '<h2 class="resoluci-n">RESOLUCIÓN SBS Nº 00123-2022</h2>'
            '<p>Lima, 12 de octubre de 2022</p><p>CONSIDERANDO:</p>'
            f'{before}<p>RESUELVE:</p><p>{article}</p>{after}'
            '<p>Regístrese, comuníquese y publíquese.</p><p>1234567-1</p>'
            '</div></div></body></html>')


def parse(text):
    return parse_resolution(text, url=URL, retrieved_at='checked')[0].iloc[0]


def test_authorization_preserves_roles_and_does_not_certify_execution():
    row = parse(html())
    assert row.resolution_date == '2022-10-12' and row.resolution_number == '00123-2022'
    assert row.event_type == 'merger_authorized' and row.effective_date == ''
    assert row.absorbing_entity_name == 'Banco Ejemplo S.A.'
    assert row.absorbed_entity_name == 'Financiera Ejemplo S.A.'
    assert row.roles_status == 'operative_article' and row.evidence_text == AUTH
    assert row.source_title == TITLE and row.source_url == API
    assert len(row.html_sha256) == 64


def test_edpyme_authorization_does_not_infer_roles_from_name_order():
    article = ('Artículo Primero.- Autorizar a Servicios Ejemplo EDPYME, la fusión por absorción '
               'con la empresa Factoring Ejemplo S.A., en los términos propuestos en su solicitud.')
    row = parse(html(article, 'Autorizan a Servicios Ejemplo EDPYME, la fusión por absorción con la empresa Factoring Ejemplo S.A.'))
    assert row.event_type == 'merger_authorized' and row.roles_status == 'unspecified'
    assert row.absorbing_entity_name == '' and row.absorbed_entity_name == ''
    assert 'en los términos' in row.effective_condition


def test_clarification_extracts_stipulated_date_not_a_new_authorization():
    row = parse(html(CLARIFY, TITLE.replace('la fusión', 'la aclaratoria de la fusión')))
    assert row.event_type == 'merger_date_clarification' and row.effective_date == '2022-01-01'
    assert row.effective_basis == 'date_in_clarification_article'
    assert row.roles_status == 'unspecified'


@pytest.mark.parametrize('article', [AUTH.replace('Autorizar', 'No autorizar'),
    AUTH.replace('Autorizar', 'Denegar'), 'Artículo Primero.- Aprobar un reglamento general de fusiones.',
    'Artículo Primero.- Según una resolución anterior: '+AUTH,
    AUTH.replace('Autorizar la fusión por absorción de', 'Autorizar a')])
def test_negative_general_historical_and_non_merger_articles_need_review(article):
    row = parse(html(article))
    assert row.extraction_status == 'needs_review' and row.event_type == '' and row.entity_name == ''


def test_recitals_later_articles_and_unrelated_titles_cannot_confirm_events():
    row = parse(html('Artículo Primero.- Autorizar apertura de oficina.', before=f'<p>{AUTH}</p>', after=f'<p>{AUTH}</p>'))
    assert row.event_type == '' and row.extraction_status == 'needs_review'
    assert parse(html(title=TITLE.replace('Banco Ejemplo', 'Banco Diferente'))).event_type == ''
    assert parse(html(CLARIFY)).event_type == ''


def test_registration_conditions_are_kept_without_date_inference():
    text = AUTH + ' La fusión entrará en vigencia cuando se inscriba la escritura pública.'
    row = parse(html(text))
    assert row.event_type == 'merger_authorized' and row.effective_date == ''
    assert 'cuando se inscriba' in row.effective_condition


@pytest.mark.parametrize('transform', [lambda x:x.replace('</html>', ''),
    lambda x:x.replace('x1234567-1', 'x7654321-1'), lambda x:x.replace('<p>1234567-1</p>', ''),
    lambda x:x.replace('RESUELVE:', 'CONSIDERANDO:'),
    lambda x:x.replace('12 de octubre', '31 de febrero'),
    lambda x:x.replace('RESOLUCIÓN SBS', 'RESOLUCIÓN MINISTERIAL'),
    lambda x:x.replace('<p>RESUELVE:</p>', '<p>RESUELVE:</p><p>RESUELVE:</p>'),
    lambda x:x.replace('</body>', '<div class="story">Otra norma</div></body>')])
def test_incomplete_wrong_or_ambiguous_documents_fail(transform):
    with pytest.raises(SchemaChangedError):parse(transform(html()))


def test_invalid_clarified_date_fails():
    with pytest.raises(SchemaChangedError):
        parse(html(CLARIFY.replace('01 de enero','31 de febrero'), TITLE.replace('la fusión','la aclaratoria de la fusión')))


@pytest.mark.parametrize('url', [URL.replace('https','http'), URL.replace('busquedas.elperuano.pe','other.test'),
    URL+'?id=1', URL+'#fragment', URL.replace('NL','XX'), URL.replace('-1','-0'), URL.replace('1234567','123')])
def test_only_unambiguous_official_resolution_urls_allowed(url):
    with pytest.raises(InvalidQueryError):resolution_reference(url)


def test_cache_reuse_redirect_failure_and_invalid_structure_preserve_data(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT', str(tmp_path))
    provider = MergerResolutionsProvider(CATALOG['pe.elperuano.fusiones'])
    calls = []
    def request(*a, **kw):
        calls.append(a)
        return SimpleNamespace(text=html(), url=API)
    monkeypatch.setattr(provider.transport, 'request', request)
    assert provider.sync(urls=[URL, API], keep_raw=True).downloaded == 1
    assert provider.sync(urls=[URL]).skipped_existing == 1 and len(calls) == 1
    previous = provider.load().copy()
    monkeypatch.setattr(provider.transport, 'request', lambda *a,**kw:SimpleNamespace(text=html(),url=API.replace('-1','-2')))
    assert provider.sync(urls=[URL], force=True).failed == 1 and provider.load().equals(previous)
    monkeypatch.setattr(provider.transport, 'request', lambda *a,**kw:SimpleNamespace(text='<html>error</html>',url=API))
    assert provider.sync(urls=[URL], force=True).failed == 1 and provider.load().equals(previous)


def test_cli_offline_report_and_corrupt_cache_rejection(tmp_path, monkeypatch):
    from fuentes_financieras.cli import fusiones as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    provider = MergerResolutionsProvider(CATALOG['pe.elperuano.fusiones'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(text=html(),url=API))
    assert provider.sync(urls=[URL]).downloaded == 1
    monkeypatch.setattr(cli, 'source', lambda _:provider)
    def no_network(*a,**kw):raise AssertionError('Unexpected offline request')
    monkeypatch.setattr(provider.transport,'request',no_network)
    args = ['--urls',URL,'--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args) == 0
    path = tmp_path/'reports/fusiones.xlsx'
    book = load_workbook(path)
    for sheet in book:
        assert sheet.freeze_panes == 'A2' and not sheet.column_dimensions
        assert next(iter(sheet.tables.values())).tableStyleInfo.name == 'TableStyleLight9'
    assert 'Fecha de vigencia estipulada; no prueba ejecución' in [c.value for c in book['resoluciones'][1]]
    path.unlink()
    provider.storage.partition_path('resoluciones').write_bytes(b'damaged')
    assert cli.main(args) == 1 and not path.exists()


def test_cli_review_only_selection_returns_partial_status(tmp_path, monkeypatch):
    from fuentes_financieras.cli import fusiones as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    provider = MergerResolutionsProvider(CATALOG['pe.elperuano.fusiones'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(text=html('Artículo Primero.- Aprobar otro asunto.'),url=API))
    assert provider.sync(urls=[URL]).downloaded == 1
    monkeypatch.setattr(cli,'source',lambda _:provider)
    assert cli.main(['--urls',URL,'--load-only','--output-dir',str(tmp_path/'reports')]) == 2
    book = load_workbook(tmp_path/'reports/fusiones.xlsx')
    assert book['disposiciones'].max_row == 1 and book['resoluciones'].max_row == 2
