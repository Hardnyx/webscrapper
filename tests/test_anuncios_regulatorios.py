"""Recognize explicit dated dispositions and preserve ambiguous announcements."""
from types import SimpleNamespace
import pytest
from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.providers.sbs.anuncios_regulatorios import (
    announcement_reference, parse_announcement, RegulatoryAnnouncementsProvider,
)
from fuentes_financieras.exceptions import InvalidQueryError, SchemaChangedError

URL = 'https://www.sbs.gob.pe/noticia/detallenoticia/idnoticia/123'
LEAD = ('Lima, 11 de julio de 2024.- En resguardo de los intereses de los ahorristas, '
        'la Superintendencia de Banca, Seguros y AFP (SBS) ha intervenido a la '
        'Caja Municipal de Ahorro y Crédito Ejemplo S.A. (CMAC Ejemplo) por haber incurrido en una causal.')


def html(lead=LEAD, title='SBS interviene CMAC Ejemplo', extra=''):
    return f'<html><body><nav>Otra entidad en liquidación</nav><h2 class="boletin-sala-prensa__title">{title}</h2><div class="boletin-sala-prensa__text"><p>{lead}</p>{extra}</div></body></html>'


def parse(text):
    return parse_announcement(text, url=URL, retrieved_at='checked')[0].iloc[0]


def test_dated_official_intervention_keeps_name_and_no_legal_date():
    row = parse(html())
    assert row.event_type == 'intervention_announced'
    assert row.entity_name == 'Caja Municipal de Ahorro y Crédito Ejemplo S.A.'
    assert row.announcement_date == '2024-07-11' and row.effective_date == ''
    assert row.resolution_number == '' and row.evidence_text == LEAD
    assert len(row.html_sha256) == 64


def test_dissolution_and_liquidation_resolution_is_explicit():
    lead = ('Lima, 11 de agosto de 2023.- Con la finalidad de cautelar el valor de los activos de la '
            'Caja Rural de Ahorro y Crédito Ejemplo S.A.A. (CRAC Ejemplo) en Intervención, '
            'la Superintendencia de Banca, Seguros y AFP (SBS), mediante Resolución N.°2672-2023, '
            'ha dispuesto su disolución y el inicio del proceso de liquidación.')
    row = parse(html(lead, 'SBS dispone la disolución e inicio de liquidación'))
    assert row.event_type == 'dissolution_liquidation_announced'
    assert row.resolution_number == '2672-2023' and row.effective_date == ''


@pytest.mark.parametrize('lead', [LEAD.replace('ha intervenido','no ha intervenido'),
    LEAD.replace('ha intervenido','podría intervenir'), LEAD.replace('Lima, 11 de julio de 2024.- ', ''),
    'Según el informe, '+LEAD, LEAD.replace('S.A. (CMAC Ejemplo)', 'sin nombre legal')])
def test_negation_hypothesis_and_undated_text_do_not_become_events(lead):
    row = parse(html(lead))
    assert row.extraction_status == 'needs_review' and row.event_type == '' and row.entity_name == ''


def test_only_title_or_later_historical_paragraph_does_not_confirm_event():
    row = parse(html('Lima, 11 de julio de 2024.- La SBS modifica un reglamento.', extra=f'<p>{LEAD}</p>'))
    assert row.extraction_status == 'needs_review' and row.event_type == ''
    row = parse(html(LEAD, 'SBS comenta una disposición diferente'))
    assert row.event_type == ''


def test_invalid_date_or_ambiguous_structure_fails():
    with pytest.raises(SchemaChangedError):parse(html(LEAD.replace('11 de julio','31 de febrero')))
    with pytest.raises(SchemaChangedError):parse(html() + '<div class="boletin-sala-prensa__text"><p>duplicado</p></div>')
    with pytest.raises(SchemaChangedError):parse('<html>Access denied</html>')
    with pytest.raises(SchemaChangedError):parse(html().replace('</body></html>', ''))


@pytest.mark.parametrize('url', [URL.replace('https','http'),URL.replace('www.sbs.gob.pe','other.test'),
    URL+'?idnoticia=456',URL+'#fragment',URL+'?title=a&title=b',URL.replace('/123','/0')])
def test_only_unambiguous_official_news_urls_are_allowed(url):
    with pytest.raises(InvalidQueryError):announcement_reference(url)


def test_cache_reuse_rejected_redirect_and_failure_preserve_capture(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    provider=RegulatoryAnnouncementsProvider(CATALOG['pe.sbs.anuncios_regulatorios']);calls=[]
    def request(*a,**kw):calls.append(a);return SimpleNamespace(text=html(),url=URL)
    monkeypatch.setattr(provider.transport,'request',request)
    assert provider.sync(urls=[URL, URL+'?title=Example'],keep_raw=True).downloaded==1
    assert provider.sync(urls=[URL]).skipped_existing==1 and len(calls)==1
    previous=provider.load().copy()
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(text=html(),url=URL.replace('/123','/456')))
    assert provider.sync(urls=[URL],force=True).failed==1 and provider.load().equals(previous)
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(text='<html>error</html>',url=URL))
    assert provider.sync(urls=[URL],force=True).failed==1 and provider.load().equals(previous)


def test_cli_partial_coverage_and_corrupt_cache_rejection(tmp_path,monkeypatch):
    from fuentes_financieras.cli import anuncios_regulatorios as cli
    from openpyxl import load_workbook
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    provider=RegulatoryAnnouncementsProvider(CATALOG['pe.sbs.anuncios_regulatorios'])
    monkeypatch.setattr(provider.transport,'request',lambda *a,**kw:SimpleNamespace(text=html('Lima, 11 de julio de 2024.- Otra noticia.'),url=URL))
    assert provider.sync(urls=[URL]).downloaded==1
    monkeypatch.setattr(cli,'source',lambda _:provider)
    def no_network(*a,**kw):raise AssertionError('Unexpected offline request')
    monkeypatch.setattr(provider.transport,'request',no_network)
    args=['--urls',URL,'--load-only','--output-dir',str(tmp_path/'reports')]
    assert cli.main(args)==2
    path=tmp_path/'reports/anuncios_regulatorios.xlsx';book=load_workbook(path)
    assert book['eventos'].max_row==1 and book['anuncios'].max_row==2
    assert book['anuncios'].freeze_panes=='A2' and not book['anuncios'].column_dimensions
    assert next(iter(book['anuncios'].tables.values())).tableStyleInfo.name=='TableStyleLight9'
    path.unlink();provider.storage.partition_path('noticias').write_bytes(b'corrupt')
    assert cli.main(args)==1 and not path.exists()
