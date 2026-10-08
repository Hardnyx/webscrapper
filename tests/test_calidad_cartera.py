"""Guard published definitions, exposure bases, missing tokens and balance signs."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from fuentes_financieras.catalog import CATALOG
from fuentes_financieras.exceptions import SchemaChangedError
from fuentes_financieras.providers.sbs.calidad_cartera import (
    CATEGORIES, EXPECTED_QUALITY, CreditRiskCategoriesProvider,
    parse_quality, parse_categories, parse_arrears,
)
from fuentes_financieras.providers.sbs.saldos_cartera import BALANCE_METRICS, NET_CREDIT, parse_balances


def save(book):
    stream = BytesIO(); book.save(stream)
    return stream.getvalue()


def vertical(kind='categories', *, period='2026-08-31', dash=False, mismatch=False, unknown=False, fmt='General'):
    book = Workbook(); s = book.active
    s.cell(1, 1, 'Estructura de Créditos Directos e Indirectos según Categoría de Riesgo' if kind == 'categories' else 'Ratios de Morosidad según días de incumplimiento')
    s.cell(2, 1, datetime.fromisoformat(period)); s.cell(3, 1, '(En porcentaje)')
    s.cell(5, 1, 'Empresas')
    if kind == 'categories':
        for col, label in enumerate(CATEGORIES, start=2):s.cell(5, col, label)
        s.cell(5, 7, 'Total Créditos Directos e Indirectos 1/ (En miles de soles)')
        values = [90, 3, 2, 2, 3, 1000]
    else:
        s.cell(5, 2, 'Porcentaje de créditos con'); s.cell(5, 6, 'Morosidad según criterio contable SBS2/')
        for col, days in enumerate((30, 60, 90, 120), start=2):s.cell(6, col, f'Más de {days} días de incumplimiento')
        values = [5, 4, 3, 2, 4.5]
    if unknown:s.cell(5 if kind == 'categories' else 6, 3, 'Encabezado inesperado')
    for row, name in ((8, 'Entidad'), (9, 'TOTAL SISTEMA')):
        s.cell(row, 1, name)
        for col, value in enumerate(values, start=2):
            s.cell(row, col, '-' if dash and col == 4 else value+(10 if mismatch and col == 4 else 0))
            s.cell(row, col).number_format=fmt
    s.cell(11, 1, 'Fuente: Anexo SBS')
    return save(book)


def quality(*, tipo='B', changed=False, repeated=True):
    book=Workbook(); s=book.active
    blocks=(1, 4) if repeated else (1,)
    for col in blocks:
        s.cell(1, col, 'Indicadores Financieros por Empresa'); s.cell(2, col, datetime(2026,8,31))
        s.cell(3, col, '(En porcentaje)'); s.cell(5, col+1, 'Entidad' if col == 1 else 'TOTAL SISTEMA')
        s.cell(7, col, 'CALIDAD DE ACTIVOS')
        labels = sorted(EXPECTED_QUALITY[tipo])
        for row,label in enumerate(labels,start=8):
            s.cell(row,col,label+(' ****' if label.startswith('Cartera') else ''))
            s.cell(row,col+1,150 if label.startswith('Provisiones') else 5)
        s.cell(8+len(labels),col,'EFICIENCIA Y GESTIÓN')
        s.cell(10+len(labels),col,'Nota: Glosario de términos')
    if changed:s.cell(8, 1, 'Ratio nuevo no revisado')
    return save(book)


def balances(*, missing=False, mismatch=False):
    book = Workbook(); s = book.active; s.title='balance'; income=book.create_sheet('income')
    def block(sheet, start, title, section, accounts):
        sheet.cell(start,1,title); sheet.cell(start+1,1,datetime(2026,8,31)); sheet.cell(start+2,1,'(En miles de soles)')
        sheet.cell(start+4,1,section); sheet.cell(start+4,2,'Entidad')
        for c,cur in enumerate(('MN','ME','TOTAL'),2):sheet.cell(start+5,c,cur)
        for row,(label,value) in enumerate(accounts,start+7):
            sheet.cell(row,1,label)
            for col,multiplier in ((2,1),(3,0),(4,1)):sheet.cell(row,col,value*multiplier)
    values=[88,100,5,10,6,4,-20,-7]
    credit=list(zip(BALANCE_METRICS,values))
    if missing:credit=[(l,v) for l,v in credit if l!='Provisiones']
    if mismatch:credit=[(l,90 if l==NET_CREDIT else v) for l,v in credit]
    accounts=[('TOTAL ACTIVO',100),('Provisiones',-999)]+credit+[('CUENTAS POR COBRAR NETAS DE PROVISIONES',12)]
    block(s,1,'Balance General por Empresa','Activo',accounts)
    block(s,35,'Balance General por Empresa','Pasivo',[('TOTAL PASIVO',80),('PATRIMONIO',20),('TOTAL PASIVO Y PATRIMONIO',100),('RESULTADO NETO DEL EJERCICIO',5)])
    block(income,1,'Estado de Ganancias por Empresa','',[('RESULTADO NETO DEL EJERCICIO',5)])
    return save(book)


def parse(content, parser=parse_categories, tipo='B'):
    return parser(content, entity_type=tipo, period='2026-08', source_url='test', retrieved_at='test')[0]


def test_categories_preserve_dash_and_credit_equivalent_total():
    data=parse(vertical(dash=True))
    assert len(data)==12 and data.value.isna().sum()==2
    assert set(data[data.value.isna()].source_value_token)=={'-'}
    assert set(data.credit_scope)=={'direct_and_credit_equivalent_indirect'}
    totals=data[data.metric=='total_classified_credit']
    assert set(totals.value)=={1000} and set(totals.unit)=={'thousands_PEN'}
    assert set(data[data.metric=='risk_category_share'].risk_category)=={'normal','potential_problems','substandard','doubtful','loss'}


@pytest.mark.parametrize('kwargs',[{'period':'2026-07-31'},{'unknown':True},{'fmt':'0%'}])
def test_categories_reject_wrong_dates_labels_or_percentage_scale(kwargs):
    with pytest.raises(SchemaChangedError):parse(vertical(**kwargs))


def test_published_category_sum_mismatch_is_preserved():
    data=parse(vertical(mismatch=True))
    assert set(data.data_quality_flags)=={'published_categories_sum_mismatch'}
    assert data[data.risk_category=='substandard'].value.tolist()==[12,12]


def test_arrears_thresholds_are_cumulative_and_accounting_criterion_is_separate():
    data=parse(vertical('arrears'),parse_arrears)
    assert set(data.arrears_threshold_days.dropna())=={30,60,90,120}
    assert data[data.metric=='past_due_ratio_sbs'].value.tolist()==[4.5,4.5]
    inconsistent=parse(vertical('arrears',mismatch=True),parse_arrears)
    assert set(inconsistent.data_quality_flags)=={'published_arrears_order_mismatch'}


def test_quality_repeated_label_blocks_and_missing_referenced_definition():
    data=parse(quality(tipo='C'),parse_quality,'C')
    assert len(data)==16 and set(data.entity_scope)=={'entity','system_aggregate'}
    adjusted=data[data.metric.str.startswith('adjusted_')]
    assert set(adjusted.data_quality_flags)=={'source_definition_missing'}
    assert set(data[data.metric=='provisions_past_due_coverage'].value)=={150}
    assert set(data.unit)=={'percent'}
    with pytest.raises(SchemaChangedError):parse(quality(changed=True),parse_quality)


def test_balances_select_only_credit_provisions_and_preserve_negative_sign():
    data=parse(balances(),parse_balances)
    assert len(data)==24
    assert set(data.unit)=={'thousands_PEN'} and set(data.currency)=={'MN','ME','TOTAL'}
    assert data[(data.metric=='credit_provisions_contra_asset') & (data.currency=='TOTAL')].value.tolist()==[-20]
    assert -999 not in data.value.tolist()
    assert set(data.data_quality_flags)=={''}
    with pytest.raises(SchemaChangedError):parse(balances(missing=True),parse_balances)
    mismatched=parse(balances(mismatch=True),parse_balances)
    assert 'published_credit_components_mismatch' in set(mismatched.data_quality_flags)


def test_cache_skips_network_and_month_without_link_is_unavailable(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCIAL_SOURCES_DATA_ROOT',str(tmp_path))
    provider=CreditRiskCategoriesProvider(CATALOG['pe.sbs.categorias_riesgo_cartera']); calls=[]
    url='https://intranet2.sbs.gob.pe/estadistica/financiera/2026/Agosto/B-2309-ag2026.XLS'
    def request(method,target):
        calls.append(target)
        return SimpleNamespace(text=f'<a href="{url}">Agosto</a>',content=vertical())
    monkeypatch.setattr(provider.transport,'request',request)
    assert provider.sync(desde='2026-08',tipos=['B']).downloaded==1
    assert provider.sync(desde='2026-08',tipos=['B']).skipped_existing==1
    assert len(calls)==2
    assert provider.sync(desde='2026-09',tipos=['B']).unavailable==1


def test_cli_does_not_export_incomplete_source(tmp_path,monkeypatch):
    from fuentes_financieras.cli import calidad_cartera
    fake=SimpleNamespace(plan_sync=lambda **kw: [],sync=lambda **kw: SimpleNamespace(failed=1,unavailable=0))
    monkeypatch.setattr(calidad_cartera,'source',lambda dataset: fake)
    assert calidad_cartera.main(['--desde','2026-08','--output-dir',str(tmp_path)])==1
    assert not (tmp_path/'calidad_cartera.xlsx').exists()
